"""Bear / base / bull cases and a forced early sale, for a deal built with build_model.py.
Run:  python3 cases.py ../deals/<deal>/deal.json [recalc.py]
Reads deal.json "cases": "plan" (owner-time plan, applied to every case), "bear", "bull" (input overrides on top
of the plan; base = plan only). Writes cases.json in the deal folder and prints a summary."""
import json, os, shutil, subprocess, sys, tempfile
from openpyxl import load_workbook

DEAL_PATH = os.path.abspath(sys.argv[1])
HERE = os.path.dirname(DEAL_PATH)
DEAL = json.load(open(DEAL_PATH))
XL = os.path.join(HERE, DEAL.get("output", "Underwriting_Model.xlsx"))
META = json.load(open(XL + ".meta.json"))
RECALC = sys.argv[2] if len(sys.argv) > 2 else "/mnt/skills/public/xlsx/scripts/recalc.py"
N = META["N"]
C = DEAL.get("cases", {})
PLAN = C.get("plan", {})
CASES = {"bear": dict(PLAN, **C.get("bear", {})), "base": dict(PLAN), "bull": dict(PLAN, **C.get("bull", {}))}
_blank = [f"{c}.{k}" for c, o in CASES.items() for k, v in o.items() if v is None]
if _blank:
    sys.exit(f"deal.json cases have blank (null) values: {_blank}")


def inputs():
    wb = load_workbook(XL)
    out = {}
    for n, d in wb.defined_names.items():
        sh, cell = d.attr_text.split("!")
        out[n] = wb[sh][cell.replace("$", "")].value
    return out


def run(over):
    wb = load_workbook(XL)
    rf = {n: (d.attr_text.split("!")[0], d.attr_text.split("!")[1].replace("$", "")) for n, d in wb.defined_names.items()}
    for k, v in over.items(): wb[rf[k][0]][rf[k][1]] = v
    t = tempfile.mkdtemp(); p = os.path.join(t, "c.xlsx"); wb.save(p)
    subprocess.run([sys.executable, RECALC, p, "90"], check=True, capture_output=True)
    v = load_workbook(p, data_only=True); pj, st = v["Projection"], v["Structures"]
    row = lambda r: [pj.cell(r, c).value for c in range(3, 3 + N + 1)]  # yr0..yrN
    out = dict(rev=row(META["rev"]), noi=row(META["ebitda"]), sde=row(META["sde"]), capex=row(META["capex_row"]),
               irr=pj[f"C{META['irr_row']}"].value,
               npvD=st[f"C{META['comp']['D']}"].value, irrD=st[f"D{META['comp']['D']}"].value,
               npvA=st[f"C{META['comp']['A']}"].value, irrA=st[f"D{META['comp']['A']}"].value)
    shutil.rmtree(t); return out


if __name__ == "__main__":
    IN = inputs()
    res = {k: run(o) for k, o in CASES.items()}
    # Forced sale: decide ~month 6, close ~month 12 (laundromats take 3-9 months to sell).
    FS = DEAL.get("forced_sale", {})
    EXIT = FS.get("multiples", {"bear": (2.0, 2.5), "base": (2.75, 3.25), "bull": (3.25, 3.75)})  # quick / likely x buyer-verified NOI
    EQUIP_RECOVERY = FS.get("equip_recovery", 0.5)
    BROKER, BROKER_MIN, SELL_LEGAL = FS.get("broker_fee", 0.10), FS.get("broker_min", 25000), FS.get("sell_legal", 7500)
    RATE = IN["RyanFed"] + IN["RyanNJ"]
    PRICE, BUY_COSTS, WC = IN["Price"], IN["ClosingCosts"], IN["WorkingCap"]
    for k, r in res.items():
        noi1, sde1 = r["noi"][1], r["sde"][1]
        e = {}
        for tag, m in zip(("quick", "likely"), EXIT[k]):
            # Machines replaced before the sale are cash out; a buyer pays back ~half via a higher price.
            capex_pre = -sum(r["capex"][:2])
            price = m * noi1 + EQUIP_RECOVERY * capex_pre
            fee = max(BROKER * price, BROKER_MIN)
            proceeds = price - fee - SELL_LEGAL
            pretax = -(PRICE + BUY_COSTS + WC) - capex_pre + sde1 + proceeds + WC
            e[tag] = dict(mult=m, capex_pre=capex_pre, price=price, fee=fee, proceeds=proceeds, pretax=pretax,
                          aftertax_approx=pretax * (1 - RATE) if pretax < 0 else pretax)
        r["exit"] = e
    json.dump({"cases": CASES, "results": res}, open(os.path.join(HERE, "cases.json"), "w"), indent=1, default=float)
    for k, r in res.items():
        irr = r["irr"]
        print(k, "rev1", round(r["rev"][1]), "noi1", round(r["noi"][1]), "noi3", round(r["noi"][3]),
              "sde1", round(r["sde"][1]), "capex", round(-sum(r["capex"])),
              "irr", round(irr, 3) if isinstance(irr, float) else irr, "npvD", round(r["npvD"]))
        for tag, e in r["exit"].items():
            print("   forced sale", tag, {a: round(b) if b > 100 else b for a, b in e.items()})
