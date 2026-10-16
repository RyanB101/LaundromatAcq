"""Structure E (family partnership, parents' capital + Ryan's profits interest, the standing structure) vs D
(Ryan's LLC, AFR loan, S-corp later). Pulls yearly pre-tax lines from the deal's model (LibreOffice recalc),
recomputes D with the model's own rules (checked against the Structures sheet), computes E and its
capital-account-safe variants E1/E2 (with and without Ryan's S corp), and a bear-case sale after Year 1.
Run:  python3 structure_e.py ../deals/<deal>/deal.json [recalc.py]   -> prints a summary, writes structure_e.json
Tax rates, parents' funding and S-corp terms come from the model's Inputs sheet; partnership terms (pref, split,
later S-corp salary) from deal.json "structure" (defaults = the standing deal terms).
"""
import json, os, shutil, subprocess, sys, tempfile
from openpyxl import load_workbook

DEAL_PATH = os.path.abspath(sys.argv[1])
HERE = os.path.dirname(DEAL_PATH)
DEAL = json.load(open(DEAL_PATH))
XL = os.path.join(HERE, DEAL.get("output", "Underwriting_Model.xlsx"))
META = json.load(open(XL + ".meta.json"))
RECALC = sys.argv[2] if len(sys.argv) > 2 else "/mnt/skills/public/xlsx/scripts/recalc.py"
N = META["N"]
_C = DEAL.get("cases", {})
PLAN = dict(_C.get("plan", {}))
CASES = {"bear": dict(PLAN, **_C.get("bear", {})), "base": dict(PLAN)}


def _inputs():
    wb = load_workbook(XL)
    out = {}
    for n, dn in wb.defined_names.items():
        sh, c = dn.attr_text.split("!")
        out[n] = wb[sh][c.replace("$", "")].value
    return out


IN = _inputs()
# Tax parameters (model inputs, plus the structure E spec)
R_FED, R_NJ, R_OTHER, EBL_S = IN["RyanFed"], IN["RyanNJ"], IN["RyanOtherInc"], IN["EBLSingle"]
P_FED, P_NIIT, P_NJ = IN["ParentFed"], IN["NIIT"], IN["ParentNJ"]
P_ORD = P_FED + P_NIIT                      # passive ordinary income/loss, federal
P_LTCG = 0.20 + P_NIIT                      # capital loss used against long-term gains
SS0, WG, QBI, EXIT_TAX, DISC = IN["SSBase"], IN["WageGrowth"], IN["QBIRate"], IN["ExitTax"], IN["DiscRate"]
LOAN, RATE, FORGIVE, SAL, SCP = IN["LoanAmt"], IN["LoanRate"], IN["LoanForgive"], IN["RyanSalary"], IN["SCorpPayroll"]
SCORP_YEAR = IN["SCorpYear"]
PRICE, CLOSE, WC = IN["Price"], IN["ClosingCosts"], IN["WorkingCap"]
TOTAL = PRICE + CLOSE + WC                  # cash at close
CUSH = LOAN - TOTAL                         # parents' funding left in the business as cash
ST = DEAL.get("structure", {})
PREF = ST.get("pref", 0.05); R_SPLIT = ST.get("ryan_split", 0.80)
SAL_LATER = ST.get("scorp_salary_later", 25000)   # Ryan's S-corp salary at ~10 hrs/wk (Year 2 dollars)


def model_rows(over):
    wb = load_workbook(XL)
    for n, dn in wb.defined_names.items():
        if n in over:
            sh, c = dn.attr_text.split("!"); wb[sh][c.replace("$", "")] = over[n]
    t = tempfile.mkdtemp(); p = os.path.join(t, "c.xlsx"); wb.save(p)
    subprocess.run([sys.executable, RECALC, p, "120"], check=True, capture_output=True)
    st = load_workbook(p, data_only=True)["Structures"]
    R = META["st"]
    row = lambda r: [st.cell(r, c).value or 0 for c in range(3, 3 + N + 1)]
    out = dict(ebitda=row(R["ebitda"]), h=row(R["ownlab"]), wages=row(R["wages"]), fdep=row(R["fdep"]), njdep=row(R["ndep"]),
               cf=row(R["cf"]), exitv=row(R["exit"]), d_family=row(R["d_fam"]), d_npv=st[f"C{META['comp']['D']}"].value)
    shutil.rmtree(t)
    return out


def npv(cfs, r=DISC):
    return sum(c / (1 + r) ** t for t, c in enumerate(cfs))


def se_tax(earn, y):
    if earn <= 0: return 0.0
    base = SS0 * (1 + WG) ** (y - 1)
    return min(earn * 0.9235, base) * 0.124 + earn * 0.9235 * 0.029


# ---------- D: same rules as the model's Structures sheet ----------
def run_D(m, last=None, scorp=True):
    last = last or N
    fam, ryan, par = [m["cf"][0]], [CUSH], [-LOAN]
    bal, fcf, ncf = LOAN, 0.0, 0.0
    for y in range(1, last + 1):
        e, h, fd, nd = m["ebitda"][y], m["h"][y], m["fdep"][y], m["njdep"][y]
        intr = bal * RATE
        s = scorp and y >= SCORP_YEAR
        sal = SAL * (1 + WG) ** max(0, y - SCORP_YEAR) if s else 0
        inc = e + h - fd - intr
        emp = sal * (SCP - 0.0765)
        se = sal * SCP if s else se_tax(inc, y)
        qbi = max(0, inc - (sal + emp if s else se / 2)) * QBI
        ti = inc - (emp if s else se / 2) - qbi
        used = min(fcf, max(0, ti)); other = min(-ti, min(R_OTHER, EBL_S)) if ti < 0 else 0
        fcf = fcf - used + max(0, -ti) - other
        fed = (max(0, ti) - used - other) * R_FED
        nji = e + h - nd - intr - emp
        nused = min(ncf, max(0, nji)); ncf = ncf - nused + max(0, -nji)
        nj = (max(0, nji) - nused) * R_NJ
        tax = fed + se + nj
        ptax = intr * (P_ORD + P_NJ)
        xt = m["exitv"][y] * EXIT_TAX
        fam.append(m["cf"][y] + h - tax - ptax - xt)
        par.append(intr - ptax)
        ryan.append(m["cf"][y] + h - tax - intr - xt)
        bal = max(0, bal - FORGIVE)
    return dict(fam=fam, ryan=ryan, par=par, npv=npv(fam), npv_ryan=npv(ryan), npv_par=npv(par), nol_left=fcf)


# ---------- E: partnership, parents' capital, Ryan profits interest ----------
def run_E(m, last=None):
    """GP to Ryan = value of his hours (SE-taxable). Cash after GP: 5% cumulative pref to parents on unreturned capital,
    then 80/20. Shortfalls (capex years) are funded by parents as added capital. Tax: all fed/NJ depreciation and any
    net loss to parents; later income first charges back parents' cumulative net losses, then follows the cash split.
    Exit (last yr): capital accounts first, then 80/20; exit tax at the model's flat rate, borne pro rata to exit cash."""
    last = last or N
    contrib = -m["cf"][0] + CUSH
    fam, ryan, par = [-contrib], [0.0], [-contrib]          # parents' funding in; cash at close spent; cushion stays as partnership cash
    cash = float(CUSH)
    cap_p, cap_r = contrib, 0.0            # book capital accounts (tax = book; cash contributions only)
    cb_f = cb_n = 0.0                      # parents' cumulative net loss not yet charged back (fed, NJ)
    pref_due = 0.0; min_cap_r = 0.0
    for y in range(1, last + 1):
        e, h, fd, nd = m["ebitda"][y], m["h"][y], m["fdep"][y], m["njdep"][y]
        capex = -(m["cf"][y] - e - m["exitv"][y] - (WC if y == N else 0))  # capex this year (positive)
        # --- cash ---
        cash += e - capex                  # GP paid to Ryan out of e + h (store saves h in wages)
        add = max(0, -cash); cash += add; contrib += add; cap_p += add
        pref_due += PREF * contrib
        dist = cash if y < last else max(0, cash)
        pp = min(dist, pref_due); pref_due -= pp; rest = dist - pp
        dp, dr = pp + rest * (1 - R_SPLIT), rest * R_SPLIT
        cash -= dist
        # --- tax allocation (fed and NJ tracked separately) ---
        def alloc(inc_before_dep, dep, cb):
            sp = -dep; sr = 0.0
            x = max(0, inc_before_dep)
            c = min(x, cb); sp += c; cb -= c; x -= c
            pi = min(x, pp); sp += pi; x -= pi
            sp += x * (1 - R_SPLIT); sr += x * R_SPLIT
            if inc_before_dep < 0: sp += inc_before_dep
            if sp < 0: cb += -sp
            return sp, sr, cb
        spf, srf, cb_f = alloc(e, fd, cb_f)
        spn, srn, cb_n = alloc(e, nd, cb_n)
        cap_p += spf - dp; cap_r += srf - dr; min_cap_r = min(min_cap_r, cap_r)
        # Ryan: GP + distributive share, SE tax on both, QBI on the share only
        se = se_tax(h + srf, y)
        qbi = max(0, srf - se / 2 * (srf / (h + srf) if h + srf > 0 else 0)) * QBI
        r_tax = max(0, h + srf - se / 2 - qbi) * R_FED + se + max(0, h + srn) * R_NJ
        # Parents: passive, losses usable each year against their passive income; QBI limited by wages/UBIA (approx.)
        pqbi = min(max(0, spf) * QBI, 0.5 * m["wages"][y] * (spf / (spf + srf) if spf + srf > 0 else 0))
        p_tax = spf * P_ORD - pqbi * P_FED + spn * P_NJ
        xt = m["exitv"][y] * EXIT_TAX
        xp = xr = 0.0
        if y == last and m["exitv"][y]:
            proceeds = m["exitv"][y] + WC + cash
            to_p = min(proceeds, max(0, cap_p)); to_r = min(proceeds - to_p, max(0, cap_r))
            left = proceeds - to_p - to_r
            xp, xr = to_p + left * (1 - R_SPLIT), to_r + left * R_SPLIT
            dp -= 0; ex_tax_p, ex_tax_r = xt * xp / proceeds, xt * xr / proceeds
        else:
            ex_tax_p = ex_tax_r = 0.0
        ryan.append(h + dr + xr - r_tax - ex_tax_r)
        par.append(dp + xp - add - p_tax - ex_tax_p)
        fam.append(ryan[-1] + par[-1])
    return dict(fam=fam, ryan=ryan, par=par, npv=npv(fam), npv_ryan=npv(ryan), npv_par=npv(par),
                min_cap_ryan=min_cap_r, cb_left=cb_f)


# ---------- E1 / E2 (capital-account-safe variants), optional S-corp for Ryan's interest ----------
def ryan_tax_on(gp, share, share_nj, y, scorp, sal_target):
    """Ryan's tax on his guaranteed payment + distributive share. scorp: interest held by an S corp that pays him a
    salary (capped at the S corp's income); payroll tax at the model's 18.3% replaces SE tax on the rest."""
    inc = gp + share
    if not scorp:
        se = se_tax(inc, y)
        qbi = max(0, share - se / 2 * (share / inc if inc > 0 else 0)) * QBI
        return max(0, inc - se / 2 - qbi) * R_FED + se + max(0, gp + share_nj) * R_NJ
    sal = max(0, min(sal_target, inc))
    emp = sal * (SCP - 0.0765)
    k1 = inc - sal - emp
    qbi = max(0, min(k1, share)) * QBI
    return max(0, sal + k1 - qbi) * R_FED + sal * SCP + max(0, gp + share_nj - emp) * R_NJ + 500  # + NJ S-corp min tax/fees


def run_V(m, mode="E1", scorp=False, last=None):
    last = last or N
    contrib = -m["cf"][0] + CUSH
    fam, ryan, par = [-contrib], [0.0], [-contrib]
    cash, held = float(CUSH), 0.0              # held = Ryan's accrued-but-unpaid distributions (E2), kept as cash
    cap_p, cap_r, unret = contrib, 0.0, contrib
    cb_f = cb_n = 0.0; pref_due = 0.0; min_cap_r = 0.0
    h10 = [0] + [m["h"][2] / (1 + WG)] + [m["h"][y] for y in range(2, N + 1)]
    out_cap_p = None
    for y in range(1, last + 1):
        e, h, fd, nd = m["ebitda"][y], m["h"][y], m["fdep"][y], m["njdep"][y]
        X = m["exitv"][y]
        capex = -(m["cf"][y] - e - X - (WC if y == N else 0))
        sal_t = SAL_LATER * (1 + WG) ** max(0, y - 2) * (h / h10[y] if h10[y] else 1)
        cash += e - capex
        add = max(0, -(cash - held)); cash += add; contrib += add; cap_p += add; unret += add
        pref_due += PREF * unret
        dist = max(0, cash - held)
        pp = min(dist, pref_due); pref_due -= pp; rest = dist - pp
        dp, dr = pp + rest * (1 - R_SPLIT), rest * R_SPLIT
        if mode == "E2":
            spf, srf, cb_f = _alloc(e, fd, cb_f, pp)
            spn, srn, cb_n = _alloc(e, nd, cb_n, pp)
            avail = max(0, cap_r + srf)
            want = dr + held
            paid = min(want, avail)
            held = want - paid; cash -= dp + paid; dr = paid
        else:  # E1 target allocations: capital accounts track the liquidation waterfall at book (= tax) value
            cash -= dp + dr
            ni_f, ni_n = e - fd, e - nd
            bv = cap_p - dp + cap_r - dr + ni_f
            first = unret + pref_due
            tgt_p = min(bv, first) + max(0, bv - first) * (1 - R_SPLIT)
            tgt_r = bv - tgt_p
            srf = tgt_r - (cap_r - dr); spf = ni_f - srf
            srn = srf; spn = ni_n - srn
        cap_p += spf - dp; cap_r += srf - dr; min_cap_r = min(min_cap_r, cap_r)
        r_tax = ryan_tax_on(h, srf, srn, y, scorp, sal_t)
        pqbi = min(max(0, spf) * QBI, 0.5 * m["wages"][y] * (spf / (spf + srf) if spf + srf > 0 else 0))
        p_tax = spf * P_ORD - pqbi * P_FED + spn * P_NJ
        xp = xr = ex_p = ex_r = 0.0
        if y == last and X:
            proceeds = X + WC + cash
            if mode == "E1":
                first = unret + pref_due
                xp = min(proceeds, first) + max(0, proceeds - first) * (1 - R_SPLIT); xr = proceeds - xp
            else:
                to_p = min(proceeds, max(0, cap_p)); to_r = min(proceeds - to_p, max(0, cap_r))
                left = proceeds - to_p - to_r
                xp, xr = to_p + left * (1 - R_SPLIT), to_r + left * R_SPLIT
            out_cap_p = xp
            xt = X * EXIT_TAX; ex_p, ex_r = xt * xp / proceeds, xt * xr / proceeds
        ryan.append(h + dr + xr - r_tax - ex_r)
        par.append(dp + xp - add - p_tax - ex_p)
        fam.append(ryan[-1] + par[-1])
    return dict(fam=fam, ryan=ryan, par=par, npv=npv(fam), npv_ryan=npv(ryan), npv_par=npv(par),
                min_cap_ryan=min_cap_r, parents_exit_cash_final=out_cap_p, ryan_held_at_exit=held)


def _alloc(inc_before_dep, dep, cb, pp):
    sp = -dep; sr = 0.0
    x = max(0, inc_before_dep)
    c = min(x, cb); sp += c; cb -= c; x -= c
    pi = min(x, pp); sp += pi; x -= pi
    sp += x * (1 - R_SPLIT); sr += x * R_SPLIT
    if inc_before_dep < 0: sp += inc_before_dep
    if sp < 0: cb += -sp
    return sp, sr, cb


def early_exit_V(m, mult, scorp):
    """Year-1 sale: no distributable profit, so in every E variant Ryan has only his GP and the whole operating + sale
    loss sits in the parents' positive capital account."""
    e, h, fd, nd = m["ebitda"][1], m["h"][1], m["fdep"][1], m["njdep"][1]
    capex = -(m["cf"][1] - e)
    price = mult * e + 0.5 * capex
    ar = price - max(0.10 * price, 25000) - 7500
    cost = PRICE + CLOSE + capex
    pretax = -TOTAL + (e + h - capex) + ar + WC
    sal_t = SAL_LATER * h / (m["h"][2] / (1 + WG))
    r = ryan_tax_on(h, 0, 0, 1, scorp, sal_t)
    p = (e - fd + ar - (cost - fd)) * P_ORD + (e - nd + ar - (cost - nd)) * P_NJ
    return dict(family=pretax - r - p, ryan=h - r, parents=pretax - h + (-p))


def early_exit_D_noscorp(m, mult):
    x = early_exit(m, mult)["D"]
    return x


# ---------- Bear early exit: run Year 1, sell at year end ----------
def early_exit(m, mult):
    e, h, fd, nd = m["ebitda"][1], m["h"][1], m["fdep"][1], m["njdep"][1]
    capex = -(m["cf"][1] - e)
    price = mult * e + 0.5 * capex
    ar = price - max(0.10 * price, 25000) - 7500
    cost = PRICE + CLOSE + capex
    gain_f, gain_n = ar - (cost - fd), ar - (cost - nd)          # (negative = loss); ordinary (1245 + 1231 nets)
    pretax_family = -TOTAL + (e + h - capex) + ar + WC
    # D: Ryan's LLC (not yet S-corp). Parents fund the shortfall with more loan; unpaid balance is discharged.
    intr = LOAN * RATE
    inc = e + h - fd - intr + gain_f
    r_fed = -min(-inc, min(R_OTHER, EBL_S)) * R_FED if inc < 0 else inc * R_FED
    nol = max(0, -inc - R_OTHER)
    r_nj = max(0, e + h - nd - intr + gain_n) * R_NJ                  # NJ business loss can't offset wages; no NJ NOL value
    short = max(0, -(CUSH + e + h - intr - capex))
    owed = LOAN + short - FORGIVE                                     # annual exclusion forgiven as a gift at year end
    unpaid = max(0, owed - (ar + WC))
    # COD on 'unpaid': excluded to the extent Ryan is insolvent, then reduces his NOL dollar for dollar (Sec. 108(b)).
    nol_after = max(0, nol - unpaid)
    p_int = intr * (P_ORD + P_NJ)
    d = dict(ryan_tax=r_fed + r_nj, nol_before=nol, unpaid_loan=unpaid, nol_after_cod=nol_after,
             parents_interest_tax=p_int,
             parents_bad_debt_ltcg=-unpaid * P_LTCG, parents_bad_debt_stcg=-unpaid * (P_FED + P_NIIT))
    d["family_ltcg"] = pretax_family - d["ryan_tax"] - p_int - d["parents_bad_debt_ltcg"]
    d["family_stcg"] = pretax_family - d["ryan_tax"] - p_int - d["parents_bad_debt_stcg"]
    # E: GP to Ryan; all depreciation + sale loss to parents (capital account stays positive), usable now.
    se = se_tax(h, 1)
    r_tax_e = max(0, h - se / 2) * R_FED + se + h * R_NJ
    p_fed = (e - fd + gain_f) * P_ORD
    p_nj = (e - nd + gain_n) * P_NJ
    e_res = dict(ryan_tax=r_tax_e, parents_fed=p_fed, parents_nj=p_nj,
                 family=pretax_family - r_tax_e - p_fed - p_nj)
    return dict(mult=mult, price=price, net_proceeds=ar, pretax_family=pretax_family, fed_sale_gain=gain_f, D=d, E=e_res)


if __name__ == "__main__":
    out = {}
    for name, over in (("base", dict(PLAN)), ("bear", dict(CASES["bear"]))):
        m = model_rows(over)
        D, E, D0 = run_D(m), run_E(m), run_D(m, scorp=False)
        out[name] = dict(model_D_npv=m["d_npv"], D=D, E=E, D_no_scorp=D0)
        print(f"   D without S-corp: {D0['npv']:,.0f} (Ryan {D0['npv_ryan']:,.0f} / parents {D0['npv_par']:,.0f})")
        for mode in ("E1", "E2"):
            for sc in (False, True):
                v = run_V(m, mode, sc); out[name][mode + ("_scorp" if sc else "")] = v
                print(f"   {mode}{' +S-corp' if sc else ''}: {v['npv']:,.0f} (Ryan {v['npv_ryan']:,.0f} / parents {v['npv_par']:,.0f});"
                      f" Ryan min cap {v['min_cap_ryan']:,.0f}; parents' exit cash final yr {v['parents_exit_cash_final'] or 0:,.0f}; Ryan held {v['ryan_held_at_exit']:,.0f}")
        print(f"{name}: model D {m['d_npv']:,.0f} | recomputed D {D['npv']:,.0f} (Ryan {D['npv_ryan']:,.0f} / parents {D['npv_par']:,.0f})"
              f" | E {E['npv']:,.0f} (Ryan {E['npv_ryan']:,.0f} / parents {E['npv_par']:,.0f}); Ryan min capital acct {E['min_cap_ryan']:,.0f}")
        if name == "bear":
            out["bear_exit"] = [early_exit(m, k) for k in (2.0, 2.5)]
            out["bear_exit_V"] = {f"{k}": {sc: early_exit_V(m, k, sc == "scorp") for sc in ("plain", "scorp")} for k in (2.0, 2.5)}
            for x in out["bear_exit"]:
                print(json.dumps(x, default=lambda v: round(v)))
    json.dump(out, open(os.path.join(HERE, "structure_e.json"), "w"), indent=1)
