"""Builds the Tumble Wash (1147 Summit Ave, Jersey City) underwriting model.

Run:  python3 build_model.py  -> writes Tumble_Wash_Underwriting_Model.xlsx next to this file.
Every number in the workbook is a formula off the Inputs sheet except the
blue hardcoded inputs, the capex schedule and the comps table (all sourced).
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.comments import Comment

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Tumble_Wash_Underwriting_Model.xlsx")
N = 17  # projection years (lease term)
FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF", size=10)
BLACK = Font(name=FONT, color="000000", size=10)
GREEN = Font(name=FONT, color="008000", size=10)
BOLD = Font(name=FONT, bold=True, size=10)
HDR = Font(name=FONT, bold=True, size=10, color="FFFFFF")
TITLE = Font(name=FONT, bold=True, size=14)
NOTE = Font(name=FONT, italic=True, size=9, color="595959")
YELLOW = PatternFill("solid", fgColor="FFFF00")
HFILL = PatternFill("solid", fgColor="1F3864")
SFILL = PatternFill("solid", fgColor="D9E1F2")
TOP = Border(top=Side(style="thin"))
USD = '$#,##0;($#,##0);"-"'
USD2 = '$#,##0.00;($#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
MULT = '0.00"x"'
NUM = '#,##0;(#,##0);"-"'

wb = Workbook()


def name(n, ref):
    wb.defined_names[n] = DefinedName(n, attr_text=ref)


def setw(ws, widths):
    for k, v in widths.items():
        ws.column_dimensions[k].width = v


def title(ws, text, sub=None):
    ws["B2"] = text
    ws["B2"].font = TITLE
    if sub:
        ws["B3"] = sub
        ws["B3"].font = NOTE


def section(ws, r, text, ncols=4):
    for c in range(2, 2 + ncols):
        ws.cell(r, c).fill = HFILL
    ws.cell(r, 2, text).font = HDR


# ---------------------------------------------------------------- Inputs
inp = wb.active
inp.title = "Inputs"
setw(inp, {"A": 2, "B": 46, "C": 14, "D": 95})
title(inp, "Inputs: Tumble Wash Laundromat, 1147 Summit Ave, Jersey City NJ",
      "Blue = hardcoded input. Yellow = key assumption / please confirm. Year 1 = first 12 months of ownership (assumed calendar 2027). Edit only column C.")
r = 5
INPUTS = []  # (section) or (name, label, value, fmt, note, key)


def S(t):
    INPUTS.append((t,))


def I(n, label, val, fmt, note, key=False):
    INPUTS.append((n, label, val, fmt, note, key))


S("Deal")
I("Price", "Purchase price", 450000, USD, "Listing ask (BizBuySell #2564518). Change to test offers.", True)
I("ClosingCosts", "Closing costs (legal, accounting/QoE, inspection, bulk-sale filing)", 20000, USD, "Estimate. Capitalized with goodwill (15-yr).")
I("WorkingCap", "Opening working capital (card float, supplies, cash reserve)", 10000, USD, "Estimate. Recovered at exit.")
I("AllocEquip", "Form 8594 allocation: machines, card system, ATM, vending (Class V)", 60000, USD, "Estimate: ~69 machines at 15 yrs old; support with a distributor appraisal. Seller will want this low (recapture).", True)
I("AllocLHI", "Form 8594 allocation: leasehold improvements (curbs, tile, plumbing troughs, venting)", 25000, USD, "Estimate. Qualified improvement property (15-yr, bonus-eligible).")
I("AllocNoncompete", "Form 8594 allocation: seller non-compete", 15000, USD, "Estimate. Sec. 197, 15-yr. Negotiate a 5-yr / 3-mile non-compete.")
S("Revenue (Year 1, annual, from broker P&L dated 10-4-26 unless noted)")
I("SelfServe", "Self-serve washer/dryer revenue", 216000, USD, "P&L $18,000/mo; listing $4,150/wk. Verify with 24 months of CCI card-system reports.", True)
I("WFLbsMonth", "Wash-dry-fold pounds per month", 8000, NUM, "P&L ~8,000 lb/mo (listing 1,846 lb/wk). Verify in Cents POS.")
I("WFPrice", "Wash-dry-fold price per lb (Year 1)", 1.25, USD2, "P&L/listing. Local range $1.00-1.25; delivery $1.25-1.49 (research/opex.md).")
I("WFPriceNew", "Wash-dry-fold price per lb after increase", 1.50, USD2, "Assumption: raise to $1.50 (CLA survey mode $1.50-1.99).", True)
I("WFPriceYear", "Year the W&F price increase takes effect", 2, NUM, "Assumption.")
I("WFVolLoss", "W&F volume lost when price rises", 0.10, PCT, "Assumption.")
I("DryClean", "Drop-off dry cleaning revenue", 6000, USD, "P&L $500/mo.")
I("ATM", "ATM & Bitcoin machine revenue", 6000, USD, "P&L $500/mo. Seller-owned ATM included. Verify.")
I("SoapVend", "Soap vending revenue", 15000, USD, "P&L $1,250/mo. NJ sales tax applies to detergent; confirm whether P&L figure is net of tax.")
I("SnackVend", "Snack/soda vending revenue", 7200, USD, "P&L $600/mo.")
I("GameVend", "Game (quarter pusher / arcade) revenue", 6000, USD, "P&L $500/mo. Unverifiable cash; verify.")
I("AncRealize", "Share of ancillary (non-laundry) revenue underwritten", 0.85, PCT, "Haircut on cash-heavy, unverifiable lines (dry cleaning, ATM, vending, games).", True)
I("SSGrowth", "Self-serve revenue growth per year (price + volume)", 0.03, PCT, "Assumption: vend price increases roughly track inflation.")
I("WFGrowth", "W&F revenue growth per year after the step-up", 0.03, PCT, "Assumption.")
I("AncGrowth", "Ancillary revenue growth per year", 0.02, PCT, "Assumption.")
I("RetoolLift", "Self-serve revenue lift from a fully new washer/dryer fleet", 0.05, PCT, "Equipment thread base case (5%).")
S("Operating expenses (Year 1)")
I("RentYr1", "Rent incl. NNN, CAM and sewer (annual)", 33096, USD, "$2,758/mo per listing; +~3%/yr; ~17 yrs left. Confirm escalation date and that sewer is truly included (P&L says so).", True)
I("RentEsc", "Rent escalation per year", 0.03, PCT, "Listing: approx. 3%/yr.")
I("PaidHours", "Paid attendant hours per week (absentee model)", 91, NUM, "P&L: 91 hrs/wk (store open 112, 21 unattended).")
I("Wage", "Attendant wage per hour (Year 1)", 17.50, USD2, "NJ minimum wage $16.48 from 1/1/2027; local attendants ~$16.50-18 (research/opex.md). P&L uses $16.50.", True)
I("WageGrowth", "Wage growth per year", 0.035, PCT, "NJ minimum wage is CPI-indexed.")
I("Burden", "Employer payroll burden (FICA, FUTA, NJ UI/DI/WF, Jersey City 1%, workers' comp, sick leave)", 0.16, PCT, "research/opex.md: ~16% all-in. P&L shows N/A for payroll tax and WC.")
I("WaterYr1", "Water (annual)", 32400, USD, "P&L $2,700/mo. JCMUA rates rise ~5%/yr (2026 $17.82 per 1,000 gal combined).")
I("GasElecYr1", "PSE&G gas + electric (annual)", 32400, USD, "P&L $2,700/mo on the 'Gas' line, electric marked 'Inc'. Our read: one combined PSE&G bill. If electric is extra, add it in the next row.", True)
I("ElecExtra", "Additional electric not in the P&L (annual)", 0, USD, "Set to ~$10-15k if 12 months of PSE&G bills show electric is billed separately.", True)
I("UtilCont", "Utility contingency (unverified bills)", 0.10, PCT, "Assumption until 12-24 months of bills are reviewed.")
I("UtilInfl", "Utility cost inflation per year", 0.04, PCT, "PSE&G/JCMUA trend (research/opex.md).")
I("UtilSave", "Utility saving with a fully new fleet", 0.20, PCT, "Equipment thread: new washers ~25% less water/energy; blended ~20%.")
I("RepairsYr1", "Repairs & maintenance, current fleet (Year 1)", 12000, USD, "P&L $5,400/yr looks owner-subsidized; equipment thread ~$8k + parts/outside techs.")
I("RepairsAging", "Repair cost growth per year while fleet ages", 0.08, PCT, "Equipment thread.")
I("RepairsNewPct", "Repairs on a new fleet, % of revenue", 0.02, PCT, "Industry benchmark for new equipment under warranty.")
I("Insurance", "Insurance: property, GL, equipment breakdown, umbrella", 7500, USD, "Estimate $5-9k (research/opex.md). P&L shows $2,880.")
I("CardFeePct", "Card processing fees, % of laundry revenue", 0.0143, PCT, "P&L $400/mo on $28k/mo laundry revenue.")
I("POS", "POS system + POS cards (CCI, Cents)", 6600, USD, "P&L $350 + $200/mo.")
I("WFSupplyLb", "W&F supplies per lb", 0.0375, USD2, "P&L $300/mo / 8,000 lb.")
I("VendCOGS", "Vending COGS, % of soap + snack revenue", 0.34, PCT, "P&L $625/mo on $1,850/mo.")
I("DCCost", "Dry cleaning cost, % of dry cleaning revenue", 0.50, PCT, "P&L $250 on $500.")
I("TeleAlarmTrash", "Phone/Wi-Fi, alarm, fire alarm, trash", 2595, USD, "P&L $120+$30+$35/mo + $375 fire alarm.")
I("MiscSupplies", "Misc. supplies", 1800, USD, "P&L $150/mo.")
I("ProFees", "Accounting, tax prep, legal", 5000, USD, "Not in P&L. Two depreciation books (federal + NJ) and a partnership return.")
I("Licenses", "Licenses, LLC annual report, vending/health permits", 1000, USD, "Not in P&L.")
I("PayrollSvc", "Payroll service", 1200, USD, "Gusto ~$49/mo + $6/person.")
I("Marketing", "Marketing (Google profile, signage, flyers, W&F promos)", 3000, USD, "Not in P&L.")
I("ContPct", "General contingency, % of revenue", 0.01, PCT, "Unbudgeted items.")
I("ExpInfl", "General expense inflation", 0.03, PCT, "Assumption.")
S("Owner time (Ryan)")
I("OwnerHours", "Hours per week Ryan works in the store (replacing paid attendants)", 40, NUM, "Assumption. 500+ hrs/yr (~10/wk) is the material-participation test. Broker: plan 10-15 hrs/wk even 'absentee'.", True)
S("Owner hand-off and delivery/commercial growth (defaults = off; see Summary scenario table)")
I("HandoffYear", "Year Ryan steps back to a management role", 2, NUM, "Ryan: full-time in Year 1, then train staff to take over.")
I("OwnerHoursLater", "Ryan's hours per week in the store after the hand-off", 40, NUM, "Default 40 = no hand-off. Ryan's plan: ~10 (management, books, marketing, contracts).", True)
I("MgrHours", "Day-manager hours per week (lead attendant who covers Ryan's shifts) from hand-off", 0, NUM, "Default 0 = off. Plan: 40.")
I("MgrPremium", "Day-manager pay premium over attendant wage, per hour", 4.50, USD2, "~$22/hr vs $17.50 attendant (research/opex.md range).")
I("DelivStart", "First year of the pickup/delivery + commercial program", 2, NUM, "Assumption.")
I("DelivLbsRamp", "Delivery/commercial pounds per month, first year", 0, NUM, "Default 0 = off. Plan: 2,000 lb/mo.", True)
I("DelivLbsFull", "Delivery/commercial pounds per month, later years", 0, NUM, "Default 0 = off. Plan: 5,000 lb/mo (e.g. ~60 residential customers + 2-4 Airbnb/gym/restaurant accounts).", True)
I("DelivPrice", "Delivery/commercial price per lb", 1.95, USD2, "Jersey City delivery $1.25-1.49 at local stores, app services ~$2.65 (research/opex.md); commercial linens taxable 6.625% (passed through).")
I("LbsPerHour", "Pounds processed (wash, dry, fold) per labor hour", 30, NUM, "Industry rule of thumb 25-35 lb/hr.")
I("DelivLaborShare", "Share of processing labor that is new paid hours (rest absorbed by idle attendant time)", 0.6, PCT, "Attendants already on shift have slack during slow hours.")
I("DelivUtilLb", "Utilities per lb processed", 0.10, USD2, "Estimate: water, gas, electric for wash + dry.")
I("DriverHours", "Driver hours per week once the program runs", 20, NUM, "Part-time driver; routes 2x/day, 5 days.")
I("DriverWage", "Driver wage per hour", 20.00, USD2, "Assumption.")
I("VehicleCost", "Van lease, fuel, commercial auto insurance per year", 12000, USD, "Assumption ~$1,000/mo. Alternative: per-stop courier.")
S("Equipment capital plan")
I("CapexScenario", "Capex scenario (1 = keep & repair, retool washers yr 8 / dryers yr 11; 2 = washers now, dryers yr 11; 3 = full retool at close)", 1, NUM, "Equipment thread recommends 1 (/mnt/project-files/equipment/README.md).", True)
I("CapexInfl", "Equipment price inflation per year (capex schedule is in today's dollars)", 0.03, PCT, "Assumption.")
I("EquipLife", "Federal/NJ MACRS recovery period for machines (5 or 7)", 5, NUM, "Asset class 57.0 = 5-yr (research/tax.md). Confirm with CPA.")
I("Bonus", "Federal bonus depreciation %", 1.0, PCT, "OBBBA: 100% permanent for property acquired after 1/19/2025; used property from unrelated seller qualifies.")
I("NJ179", "NJ Section 179 limit per year", 25000, USD, "NJ decoupled: $25k cap, no bonus (GIT-DEP).")
S("Tax rates and limits")
I("RyanFed", "Ryan: federal marginal income tax rate", 0.22, PCT, "Placeholder. Depends on Ryan's other income. Please confirm.", True)
I("RyanNJ", "Ryan: NJ marginal rate", 0.0637, PCT, "NJ GIT bracket $75k-$500k (single).")
I("RyanOtherInc", "Ryan: other income available to absorb business losses (federal)", 50000, USD, "Structure memo: ~$50k role at the parents' practice. Losses above this become an NOL carryforward.", True)
I("EBLSingle", "Excess business loss limit, single (2026)", 256000, USD, "Rev. Proc. 2025-32.")
I("SSBase", "Social Security wage base (2026)", 184500, USD, "SSA. Grown with wages in the model.")
I("QBIRate", "QBI deduction rate (Sec. 199A)", 0.20, PCT, "Permanent under OBBBA.")
I("ParentFed", "Parents: federal marginal rate", 0.37, PCT, "Assumed top bracket (physicians). Confirm.", True)
I("NIIT", "Net investment income tax on passive income", 0.038, PCT, "Sec. 1411; applies to passive owners.")
I("ParentNJ", "Parents: NJ marginal rate", 0.0897, PCT, "NJ GIT $500k-$1M bracket; 10.75% above $1M.")
I("ParentPassive", "Parents: other net passive income each year (structure C only)", 250000, USD, "Placeholder, e.g. rental K-1s from their investment LLCs. Structure B assumes $0.", True)
I("ParentNJBiz", "Parents: other NJ partnership/business income for netting (structure C only)", 250000, USD, "Placeholder: practice/investment partnership K-1s (NJ-BUS-2 netting).")
I("ExitTax", "Blended tax rate on exit proceeds", 0.30, PCT, "Mix of Sec. 1245 recapture (ordinary) and capital gain.")
S("Family loan + S-corp (structure D, ownership thread's recommendation)")
I("LoanAmt", "Intra-family loan from parents to Ryan's LLC", 500000, USD, "Structure memo: ~$500k AFR note, interest-only.")
I("LoanRate", "Loan interest rate (AFR)", 0.0461, PCT, "Oct 2026 mid-term AFR 4.61% (Rev. Rul. 2026-19); long-term 5.22%.")
I("LoanForgive", "Principal forgiven each year (annual-exclusion gifts)", 38000, USD, "$19k x 2 parents (2026 exclusion).")
I("SCorpYear", "First year taxed as an S-corp", 2, NUM, "Structure memo: elect from Year 2.")
I("RyanSalary", "Ryan's reasonable salary from the S-corp (Year of election)", 60000, USD, "Assumption; grows with wages. Must be defensible for ~40 hrs/wk management + attendant work.", True)
I("SCorpPayroll", "Payroll taxes on Ryan's salary (FICA both halves + NJ UI/DI/WF + JC 1%)", 0.183, PCT, "15.3% + ~3.0% (research/opex.md); JC 1% does not apply if Ryan lives in Jersey City.")
S("Valuation")
I("DiscRate", "Discount rate (after-tax, all-cash equity)", 0.10, PCT, "Small-business equity; parents' cash alternative ~4-5% in Treasuries.", True)
I("PreTaxTarget", "Target pre-tax unlevered return for max-price test", 0.15, PCT, "Assumption.")
I("ExitMult", "Exit value at end of Year 17, multiple of Year-17 EBITDA", 1.5, MULT, "Lease ends in Year 17; value depends on a renewal. 0 = walk away.", True)

for item in INPUTS:
    if len(item) == 1:
        r += 1
        section(inp, r, item[0], 3)
        r += 1
        continue
    n, label, val, fmt, note, key = item
    inp.cell(r, 2, label).font = BLACK
    c = inp.cell(r, 3, val)
    c.font = BLUE
    c.number_format = fmt
    if key:
        c.fill = YELLOW
    inp.cell(r, 4, note).font = NOTE
    name(n, f"Inputs!$C${r}")
    r += 1

# derived
r += 1
section(inp, r, "Derived (formulas)", 3)
r += 1
DER = [
    ("Goodwill", "Goodwill + going concern (Class VII) incl. closing costs", "=Price+ClosingCosts-AllocEquip-AllocLHI-AllocNoncompete", USD),
    ("TotalInvest", "Total cash at close (price + closing + working capital)", "=Price+ClosingCosts+WorkingCap", USD),
    ("AncYr1", "Ancillary revenue underwritten (Year 1)", "=(DryClean+ATM+SoapVend+SnackVend+GameVend)*AncRealize", USD),
]
for n, label, f, fmt in DER:
    inp.cell(r, 2, label).font = BLACK
    c = inp.cell(r, 3, f)
    c.font = BLACK
    c.number_format = fmt
    name(n, f"Inputs!$C${r}")
    r += 1

# ---------------------------------------------------------------- Capex & depreciation
dep = wb.create_sheet("Depreciation")
setw(dep, {"A": 2, "B": 44})
for y in range(0, N + 1):
    dep.column_dimensions[L(3 + y)].width = 11
title(dep, "Capex plan and depreciation (federal vs. New Jersey)",
      "Columns are years 0 (closing) to 17. Year-0 spend is placed in service in tax Year 1.")


def ycol(y):
    return L(3 + y)


R = {}
r = 5
dep.cell(r, 2, "Year").font = BOLD
for y in range(0, N + 1):
    c = dep.cell(r, 3 + y, y)
    c.number_format = "0"
    c.font = BOLD
    c.alignment = Alignment(horizontal="center")
R["year"] = r
r += 2
section(dep, r, "Capex schedule by scenario (today's $; source: equipment thread capex_schedule.csv)", 2 + N + 1)
r += 1
capex_csv = {1: {8: 372500, 11: 171000}, 2: {0: 363188, 11: 171000}, 3: {0: 529912}}
labels = {1: "1  Keep & repair; washers yr 8, dryers yr 11", 2: "2  Washers at close; dryers yr 11", 3: "3  Full retool at close"}
first_cap = r
for s in (1, 2, 3):
    dep.cell(r, 2, labels[s]).font = BLACK
    for y in range(0, N + 1):
        c = dep.cell(r, 3 + y, capex_csv[s].get(y, 0))
        c.font = BLUE
        c.number_format = USD
    r += 1
dep.cell(first_cap, 3 + N + 1, "Source: /mnt/project-files/equipment/capex_schedule.csv (2026-10-16)").font = NOTE
R["capex_sel"] = r
dep.cell(r, 2, "Selected scenario capex (inflated)").font = BOLD
for y in range(0, N + 1):
    col = ycol(y)
    c = dep.cell(r, 3 + y, f"=INDEX({col}{first_cap}:{col}{first_cap+2},CapexScenario)*(1+CapexInfl)^{col}${R['year']}")
    c.number_format = USD
r += 1
R["frac"] = r
dep.cell(r, 2, "Share of fleet that is new (drives revenue lift, utilities, repairs)").font = BLACK
tot = f"SUM($C${R['capex_sel']}:${ycol(N)}${R['capex_sel']})"
for y in range(0, N + 1):
    if y == 0:
        f = "=0"
    else:
        f = f"=IF({tot}=0,0,SUM($C${R['capex_sel']}:{ycol(y-1)}${R['capex_sel']})/{tot})"
    c = dep.cell(r, 3 + y, f)
    c.number_format = PCT
r += 2

section(dep, r, "MACRS rate tables (half-year convention), by recovery year 1-20", 2 + N + 1)
r += 1
macrs5 = [0.20, 0.32, 0.192, 0.1152, 0.1152, 0.0576]
macrs7 = [0.1429, 0.2449, 0.1749, 0.1249, 0.0893, 0.0892, 0.0893, 0.0446]
sl15 = [1 / 30] + [1 / 15] * 14 + [1 / 30]
for lbl, tbl, key in (("5-yr MACRS 200DB", macrs5, "m5"), ("7-yr MACRS 200DB", macrs7, "m7"), ("15-yr straight line (QIP)", sl15, "sl15")):
    dep.cell(r, 2, lbl).font = BLACK
    for i in range(20):
        c = dep.cell(r, 3 + i, tbl[i] if i < len(tbl) else 0)
        c.font = BLUE
        c.number_format = "0.00%"
    R[key] = r
    r += 1
R["msel"] = r
dep.cell(r, 2, "Machines: selected table").font = BLACK
for i in range(20):
    col = L(3 + i)
    c = dep.cell(r, 3 + i, f"=IF(EquipLife=7,{col}{R['m7']},{col}{R['m5']})")
    c.number_format = "0.00%"
dep.cell(R["m5"], 24, "IRS Pub 946 Table A-1").font = NOTE
r += 2

section(dep, r, "Machines placed in service by tax year", 2 + N + 1)
r += 1
R["placed"] = r
dep.cell(r, 2, "Machines placed in service (acquired fleet + capex)").font = BLACK
for y in range(1, N + 1):
    if y == 1:
        f = f"=AllocEquip+C{R['capex_sel']}+D{R['capex_sel']}"
    else:
        f = f"={ycol(y)}{R['capex_sel']}"
    c = dep.cell(r, 3 + y, f)
    c.number_format = USD
r += 2


def vintage_block(label, kind):
    """kind: 'fed' or 'nj'. returns total row"""
    global r
    section(dep, r, label, 2 + N + 1)
    r += 1
    start = r
    mrange = f"$C${R['msel']}:$V${R['msel']}"
    for v in range(1, N + 1):
        dep.cell(r, 2, f"Vintage year {v}").font = NOTE
        P = f"${ycol(v)}${R['placed']}"
        for y in range(1, N + 1):
            if y < v:
                f = "=0"
            elif kind == "fed":
                f = f"={P}*(IF({y}={v},Bonus,0)+(1-Bonus)*INDEX({mrange},{y - v + 1}))"
            else:
                f = f"=IF({y}={v},MIN(NJ179,{P}),0)+({P}-MIN(NJ179,{P}))*INDEX({mrange},{y - v + 1})"
            c = dep.cell(r, 3 + y, f)
            c.number_format = USD
            c.font = Font(name=FONT, size=8)
        r += 1
    tr = r
    dep.cell(r, 2, f"Total machine depreciation ({kind.upper()})").font = BOLD
    for y in range(1, N + 1):
        col = ycol(y)
        c = dep.cell(r, 3 + y, f"=SUM({col}{start}:{col}{r-1})")
        c.number_format = USD
        c.font = BOLD
    r += 2
    return tr


R["fed_mach"] = vintage_block("Federal: machines (bonus + MACRS)", "fed")
R["nj_mach"] = vintage_block("New Jersey: machines ($25k Sec. 179 + MACRS, no bonus)", "nj")

section(dep, r, "Leasehold improvements and Sec. 197 intangibles", 2 + N + 1)
r += 1
R["fed_lhi"] = r
dep.cell(r, 2, "Federal: leasehold improvements (QIP, bonus)").font = BLACK
for y in range(1, N + 1):
    f = f"=AllocLHI*(IF({y}=1,Bonus,0)+(1-Bonus)*INDEX($C${R['sl15']}:$V${R['sl15']},{y}))"
    dep.cell(r, 3 + y, f).number_format = USD
r += 1
R["nj_lhi"] = r
dep.cell(r, 2, "NJ: leasehold improvements (15-yr SL)").font = BLACK
for y in range(1, N + 1):
    dep.cell(r, 3 + y, f"=AllocLHI*INDEX($C${R['sl15']}:$V${R['sl15']},{y})").number_format = USD
r += 1
R["amort"] = r
dep.cell(r, 2, "Goodwill + non-compete + closing costs (15-yr, both)").font = BLACK
for y in range(1, N + 1):
    dep.cell(r, 3 + y, f"=IF({y}<=15,(Goodwill+AllocNoncompete)/15,0)").number_format = USD
r += 2
R["fed_dep"] = r
dep.cell(r, 2, "TOTAL FEDERAL depreciation & amortization").font = BOLD
R["nj_dep"] = r + 1
dep.cell(r + 1, 2, "TOTAL NJ depreciation & amortization").font = BOLD
R["ubia"] = r + 2
dep.cell(r + 2, 2, "Cumulative depreciable property (UBIA proxy for QBI wage/property test)").font = BLACK
for y in range(1, N + 1):
    col = ycol(y)
    for rr, f in ((r, f"={col}{R['fed_mach']}+{col}{R['fed_lhi']}+{col}{R['amort']}"),
                  (r + 1, f"={col}{R['nj_mach']}+{col}{R['nj_lhi']}+{col}{R['amort']}"),
                  (r + 2, f"=SUM($D${R['placed']}:{col}${R['placed']})+AllocLHI")):
        c = dep.cell(rr, 3 + y, f)
        c.number_format = USD
        c.font = BOLD if rr < r + 2 else BLACK
r += 4
dep.cell(r, 2, "Check: federal total over 17 yrs vs. NJ total (should match once fully recovered)").font = NOTE
dep.cell(r, 3, f"=SUM(D{R['fed_dep']}:{ycol(N)}{R['fed_dep']})-SUM(D{R['nj_dep']}:{ycol(N)}{R['nj_dep']})").number_format = USD
dep.freeze_panes = "C6"

# ---------------------------------------------------------------- Projection
pj = wb.create_sheet("Projection", 1)
setw(pj, {"A": 2, "B": 46})
for y in range(0, N + 1):
    pj.column_dimensions[L(3 + y)].width = 11
pj.column_dimensions[L(3 + N + 1)].width = 13
title(pj, "17-year projection (pre-tax, absentee basis: every staffed hour is paid)",
      "Year 0 = closing. Taxes by ownership structure are on the Structures sheet.")
P = {}
r = 5
pj.cell(r, 2, "Year").font = BOLD
for y in range(0, N + 1):
    c = pj.cell(r, 3 + y, y)
    c.number_format = "0"
    c.font = BOLD
    c.alignment = Alignment(horizontal="center")
pj.cell(r, 3 + N + 1, "Total 1-17").font = BOLD
P["year"] = r
r += 1
pj.cell(r, 2, "Calendar year (assumed)").font = NOTE
for y in range(0, N + 1):
    pj.cell(r, 3 + y, str(2026 + y)).font = NOTE
r += 1
P["frac"] = r
pj.cell(r, 2, "Share of fleet new").font = BLACK
for y in range(0, N + 1):
    c = pj.cell(r, 3 + y, f"=Depreciation!{ycol(y)}{R['frac']}")
    c.font = GREEN
    c.number_format = PCT
r += 1
P["e"] = r  # exponent helper: y-1
pj.cell(r, 2, "Years of growth (y-1)").font = NOTE
for y in range(1, N + 1):
    pj.cell(r, 3 + y, f"={ycol(y)}{P['year']}-1").font = NOTE
r += 2


def row(label, key, fy, fmt=USD, bold=False, total=True, y0=None):
    """fy(col, y) -> formula for years 1..N"""
    global r
    pj.cell(r, 2, label).font = BOLD if bold else BLACK
    if y0 is not None:
        c = pj.cell(r, 3, y0)
        c.number_format = fmt
    for y in range(1, N + 1):
        c = pj.cell(r, 3 + y, fy(ycol(y), y))
        c.number_format = fmt
        if bold:
            c.font = BOLD
    if total:
        c = pj.cell(r, 3 + N + 1, f"=SUM(D{r}:{ycol(N)}{r})")
        c.number_format = fmt
        c.font = BOLD
    P[key] = r
    r += 1
    return r - 1


section(pj, r, "Revenue", 2 + N + 2)
r += 1
e = lambda c: f"{c}${P['e']}"
fr = lambda c: f"{c}${P['frac']}"
row("Self-serve wash & dry", "ss", lambda c, y: f"=SelfServe*(1+SSGrowth)^{e(c)}*(1+RetoolLift*{fr(c)})")
row("W&F pounds", "lbs", lambda c, y: f"=WFLbsMonth*12*IF({c}${P['year']}>=WFPriceYear,1-WFVolLoss,1)", NUM, total=True)
row("W&F price per lb", "wfp", lambda c, y: f"=IF({c}${P['year']}>=WFPriceYear,WFPriceNew,WFPrice)*(1+WFGrowth)^MAX(0,{c}${P['year']}-WFPriceYear)", USD2, total=False)
row("Wash-dry-fold (incl. pickup & delivery)", "wf", lambda c, y: f"={c}{P['lbs']}*{c}{P['wfp']}")
row("Ancillary (dry cleaning, ATM, vending, games), after haircut", "anc", lambda c, y: f"=AncYr1*(1+AncGrowth)^{e(c)}")
row("Delivery/commercial pounds", "dlbs", lambda c, y: f"=IF({c}${P['year']}<DelivStart,0,IF({c}${P['year']}=DelivStart,DelivLbsRamp,DelivLbsFull))*12", NUM)
row("Pickup/delivery + commercial wash-dry-fold", "deliv", lambda c, y: f"={c}{P['dlbs']}*DelivPrice*(1+WFGrowth)^MAX(0,{c}${P['year']}-DelivStart)")
row("Total revenue", "rev", lambda c, y: f"={c}{P['ss']}+{c}{P['wf']}+{c}{P['anc']}+{c}{P['deliv']}", bold=True)
r += 1
section(pj, r, "Operating expenses", 2 + N + 2)
r += 1
infl = lambda c: f"(1+ExpInfl)^{e(c)}"
opex_start = r
row("Rent (incl. NNN, CAM, sewer)", "rent", lambda c, y: f"=RentYr1*(1+RentEsc)^{e(c)}")
row("Attendant wages", "wages", lambda c, y: f"=PaidHours*52*Wage*(1+WageGrowth)^{e(c)}")
row("Employer payroll taxes, workers' comp, sick leave", "burden", lambda c, y: f"={c}{P['wages']}*Burden")
row("Water", "water", lambda c, y: f"=WaterYr1*(1+UtilCont)*(1+UtilInfl)^{e(c)}*(1-UtilSave*{fr(c)})")
row("PSE&G gas + electric", "gas", lambda c, y: f"=(GasElecYr1+ElecExtra)*(1+UtilCont)*(1+UtilInfl)^{e(c)}*(1-UtilSave*{fr(c)})")
row("Repairs & maintenance", "rep", lambda c, y: f"=RepairsYr1*(1+RepairsAging)^{e(c)}*(1-{fr(c)})+{c}{P['rev']}*RepairsNewPct*{fr(c)}")
row("Insurance", "ins", lambda c, y: f"=Insurance*{infl(c)}")
row("Card processing", "card", lambda c, y: f"=({c}{P['ss']}+{c}{P['wf']})*CardFeePct")
row("POS system & cards", "pos", lambda c, y: f"=POS*{infl(c)}")
row("W&F supplies", "wfs", lambda c, y: f"={c}{P['lbs']}*WFSupplyLb*{infl(c)}")
row("Vending COGS + dry cleaning cost", "cogs", lambda c, y: f"={c}{P['anc']}*((SoapVend+SnackVend)*VendCOGS+DryClean*DCCost)/(DryClean+ATM+SoapVend+SnackVend+GameVend)")
row("Phone, alarm, trash", "tel", lambda c, y: f"=TeleAlarmTrash*{infl(c)}")
row("Misc. supplies", "misc", lambda c, y: f"=MiscSupplies*{infl(c)}")
row("Accounting, tax & legal", "pro", lambda c, y: f"=ProFees*{infl(c)}")
row("Licenses & fees", "lic", lambda c, y: f"=Licenses*{infl(c)}")
row("Payroll service", "psvc", lambda c, y: f"=PayrollSvc*{infl(c)}")
row("Marketing", "mkt", lambda c, y: f"=Marketing*{infl(c)}")
row("Day-manager pay premium (after hand-off)", "mgr", lambda c, y: f"=IF({c}${P['year']}>=HandoffYear,MgrHours*52*MgrPremium*(1+WageGrowth)^{e(c)}*(1+Burden),0)")
row("Delivery program: processing labor", "dlab", lambda c, y: f"={c}{P['dlbs']}/LbsPerHour*DelivLaborShare*Wage*(1+WageGrowth)^{e(c)}*(1+Burden)")
row("Delivery program: supplies + utilities", "dsup", lambda c, y: f"={c}{P['dlbs']}*(WFSupplyLb+DelivUtilLb)*{infl(c)}")
row("Delivery program: driver + vehicle", "drv", lambda c, y: f"=IF({c}{P['dlbs']}>0,DriverHours*52*DriverWage*(1+WageGrowth)^{e(c)}*(1+Burden)+VehicleCost*{infl(c)},0)")
row("Contingency", "cont", lambda c, y: f"={c}{P['rev']}*ContPct")
opex_end = r - 1
row("Total operating expenses", "opex", lambda c, y: f"=SUM({c}{opex_start}:{c}{opex_end})", bold=True)
r += 1
row("EBITDA / net operating income (absentee)", "ebitda", lambda c, y: f"={c}{P['rev']}-{c}{P['opex']}", bold=True)
row("EBITDA margin", "margin", lambda c, y: f"=IF({c}{P['rev']}=0,0,{c}{P['ebitda']}/{c}{P['rev']})", PCT, total=False)
row("Ryan's hours per week in the store", "ownhrs", lambda c, y: f"=IF({c}${P['year']}<HandoffYear,OwnerHours,OwnerHoursLater)", NUM, total=False)
row("Value of Ryan's hours (wages + burden he replaces)", "ownlab", lambda c, y: f"=MIN({c}{P['ownhrs']},PaidHours)*52*Wage*(1+WageGrowth)^{e(c)}*(1+Burden)")
row("Owner-operator cash flow (SDE) = EBITDA + Ryan's hours", "sde", lambda c, y: f"={c}{P['ebitda']}+{c}{P['ownlab']}", bold=True)
r += 1
section(pj, r, "Capital and cash flow (pre-tax, unlevered)", 2 + N + 2)
r += 1
row("Purchase price + closing costs + working capital", "acq", lambda c, y: "=0", y0="=-TotalInvest")
row("Equipment capex (selected scenario)", "capex", lambda c, y: f"=-Depreciation!{c}{R['capex_sel']}", y0=f"=-Depreciation!C{R['capex_sel']}")
row("Exit value (Year 17) + working capital returned", "exit", lambda c, y: f"=IF({c}${P['year']}={N},ExitMult*{c}{P['ebitda']}+WorkingCap,0)")
row("Pre-tax cash flow (absentee)", "cf", lambda c, y: f"={c}{P['ebitda']}+{c}{P['acq']}+{c}{P['capex']}+{c}{P['exit']}", bold=True,
    y0=f"=C{P['acq']}+C{P['capex']}")
row("Cumulative pre-tax cash flow", "cum", lambda c, y: f"={L(3 + y - 1)}{r}+{c}{P['cf']}", total=False, y0=f"=C{P['cf']}")
r += 1
pj.cell(r, 2, "Pre-tax unlevered IRR (absentee)").font = BOLD
c = pj.cell(r, 3, f"=IRR(C{P['cf']}:{ycol(N)}{P['cf']})")
c.number_format = PCT
P["irr"] = r
r += 1
pj.cell(r, 2, "Year-1 cash-on-cash (EBITDA less Year-1 capex / total cash at close)").font = BOLD
c = pj.cell(r, 3, f"=(D{P['ebitda']}+D{P['capex']})/(TotalInvest-C{P['capex']})")
c.number_format = PCT
P["coc"] = r
r += 1
pj.cell(r, 2, "Average annual capex over 17 years (reserve equivalent)").font = BOLD
c = pj.cell(r, 3, f"=-SUM(C{P['capex']}:{ycol(N)}{P['capex']})/{N}")
c.number_format = USD
P["avgcapex"] = r
pj.freeze_panes = "C6"

# ---------------------------------------------------------------- Structures
st = wb.create_sheet("Structures", 2)
setw(st, {"A": 2, "B": 52})
for y in range(0, N + 1):
    st.column_dimensions[L(3 + y)].width = 11
st.column_dimensions[L(3 + N + 1)].width = 13
title(st, "After-tax cash flow by ownership structure (family level)",
      "A: Ryan owns (parents gift the cash), works 40 hrs/wk unpaid. D: Ryan owns, funded by an AFR loan, S-corp from Year 2. B: parents' LLC owns, parents passive, no other passive income; Ryan paid W-2 for his hours. "
      "C: same as B but parents have passive income and NJ partnership income to absorb losses. Structure advice is in the ownership thread.")
r = 5
st.cell(r, 2, "Year").font = BOLD
for y in range(0, N + 1):
    st.cell(r, 3 + y, y).font = BOLD
    st.cell(r, 3 + y).number_format = "0"
st.cell(r, 3 + N + 1, "Total 1-17").font = BOLD
SY = r
r += 1
T = {}


def srow(label, fy, fmt=USD, bold=False, y0=None, total=True, font=None):
    global r
    st.cell(r, 2, label).font = BOLD if bold else BLACK
    if y0 is not None:
        c = st.cell(r, 3, y0)
        c.number_format = fmt
        if font:
            c.font = font
    for y in range(1, N + 1):
        c = st.cell(r, 3 + y, fy(ycol(y), y))
        c.number_format = fmt
        c.font = font or (BOLD if bold else BLACK)
    if total:
        c = st.cell(r, 3 + N + 1, f"=SUM(D{r}:{ycol(N)}{r})")
        c.number_format = fmt
        c.font = BOLD
    r += 1
    return r - 1


section(st, r, "Common lines (from Projection and Depreciation)", 2 + N + 2)
r += 1
T["ebitda"] = srow("EBITDA (absentee)", lambda c, y: f"=Projection!{c}{P['ebitda']}", font=GREEN)
T["ownlab"] = srow("Value of Ryan's hours", lambda c, y: f"=Projection!{c}{P['ownlab']}", font=GREEN)
T["wages"] = srow("W-2 wages paid (all attendant hours)", lambda c, y: f"=Projection!{c}{P['wages']}", font=GREEN)
T["fdep"] = srow("Federal depreciation & amortization", lambda c, y: f"=Depreciation!{c}{R['fed_dep']}", font=GREEN)
T["ndep"] = srow("NJ depreciation & amortization", lambda c, y: f"=Depreciation!{c}{R['nj_dep']}", font=GREEN)
T["ubia"] = srow("UBIA proxy", lambda c, y: f"=Depreciation!{c}{R['ubia']}", font=GREEN, total=False)
T["cf"] = srow("Pre-tax cash flow (absentee)", lambda c, y: f"=Projection!{c}{P['cf']}", font=GREEN, y0=f"=Projection!C{P['cf']}")
T["exit"] = srow("Exit value in cash flow", lambda c, y: f"=Projection!{c}{P['exit']}-IF({c}${SY}={N},WorkingCap,0)", font=GREEN)
r += 1


def carry_block(tag, ti_row, cap_expr, rate_expr):
    """Loss handling: income uses carryforward; losses offset outside income up to cap, rest carried."""
    rows = {}
    rows["cfb"] = srow(f"{tag} loss carryforward, start of year", lambda c, y: "=0" if y == 1 else f"={L(3 + y - 1)}{r + 3}", total=False)
    rows["used"] = srow(f"{tag} carryforward used", lambda c, y: f"=MIN({c}{rows['cfb']},MAX(0,{c}{ti_row}))")
    rows["abs"] = srow(f"{tag} loss used against other income", lambda c, y: f"=IF({c}{ti_row}<0,MIN(-{c}{ti_row},{cap_expr}),0)")
    rows["cfe"] = srow(f"{tag} loss carryforward, end of year", lambda c, y: f"={c}{rows['cfb']}-{c}{rows['used']}+MAX(0,-{c}{ti_row})-{c}{rows['abs']}", total=False)
    rows["tax"] = srow(f"{tag} tax (negative = tax saved on other income)",
                       lambda c, y: f"=(MAX(0,{c}{ti_row})-{c}{rows['used']}-{c}{rows['abs']})*{rate_expr}")
    return rows


blocks = {}
# ---- Structure A
section(st, r, "A. Ryan owns 100% (single-member LLC, materially participates)", 2 + N + 2)
r += 1
a = {}
a["inc"] = srow("Net business income (EBITDA + Ryan's unpaid hours - federal depreciation)", lambda c, y: f"={c}{T['ebitda']}+{c}{T['ownlab']}-{c}{T['fdep']}")
a["se"] = srow("Self-employment tax", lambda c, y: f"=IF({c}{a['inc']}>0,MIN({c}{a['inc']}*0.9235,SSBase*(1+WageGrowth)^({c}${SY}-1))*0.124+{c}{a['inc']}*0.9235*0.029,0)")
a["qbi"] = srow("QBI deduction (below threshold: 20%)", lambda c, y: f"=MAX(0,{c}{a['inc']}-{c}{a['se']}/2)*QBIRate")
a["ti"] = srow("Federal taxable income from the business", lambda c, y: f"={c}{a['inc']}-{c}{a['se']}/2-{c}{a['qbi']}")
fa = carry_block("Fed", a["ti"], "MIN(RyanOtherInc,EBLSingle)", "RyanFed")
a["njti"] = srow("NJ taxable income (EBITDA + Ryan's hours - NJ depreciation)", lambda c, y: f"={c}{T['ebitda']}+{c}{T['ownlab']}-{c}{T['ndep']}")
na = carry_block("NJ", a["njti"], "0", "RyanNJ")
a["tax"] = srow("Total tax (federal + SE + NJ)", lambda c, y: f"={c}{fa['tax']}+{c}{a['se']}+{c}{na['tax']}", bold=True)
a["fam"] = srow("Family after-tax cash flow", lambda c, y: f"={c}{T['cf']}+{c}{T['ownlab']}-{c}{a['tax']}-{c}{T['exit']}*ExitTax", bold=True, y0=f"=C{T['cf']}")
blocks["A"] = a
r += 1


def parent_block(letter, desc, pass_cap, nj_cap):
    global r
    section(st, r, desc, 2 + N + 2)
    r += 1
    b = {}
    b["inc"] = srow("Net business income (EBITDA - federal depreciation)", lambda c, y: f"={c}{T['ebitda']}-{c}{T['fdep']}")
    b["qbi"] = srow("QBI deduction (above threshold: limited by W-2 wages / property)",
                    lambda c, y: f"=MIN(MAX(0,{c}{b['inc']})*QBIRate,MAX(0.5*{c}{T['wages']},0.25*{c}{T['wages']}+0.025*{c}{T['ubia']}))")
    b["ti"] = srow("Federal taxable income from the business", lambda c, y: f"={c}{b['inc']}-{c}{b['qbi']}")
    fb = carry_block("Fed (passive)", b["ti"], pass_cap, "(ParentFed+NIIT)")
    b["njti"] = srow("NJ taxable income (EBITDA - NJ depreciation)", lambda c, y: f"={c}{T['ebitda']}-{c}{T['ndep']}")
    nb = carry_block("NJ", b["njti"], nj_cap, "ParentNJ")
    b["tax"] = srow("Total tax (parents)", lambda c, y: f"={c}{fb['tax']}+{c}{nb['tax']}", bold=True)
    b["ryan"] = srow("Ryan's after-tax wages for his hours (stays in the family)",
                     lambda c, y: f"={c}{T['ownlab']}/(1+Burden)*(1-RyanFed-RyanNJ-0.0765)")
    b["fam"] = srow("Family after-tax cash flow", lambda c, y: f"={c}{T['cf']}-{c}{b['tax']}+{c}{b['ryan']}-{c}{T['exit']}*ExitTax", bold=True, y0=f"=C{T['cf']}")
    r += 1
    return b


blocks["B"] = parent_block("B", "B. Parents' LLC owns 100%; parents passive with no other passive income; Ryan is a paid W-2 manager", "0", "0")
blocks["C"] = parent_block("C", "C. Parents' LLC owns 100%; parents passive but have passive + NJ partnership income to absorb losses", "ParentPassive", "ParentNJBiz")

# ---- Structure D
section(st, r, "D. Ryan's LLC owns 100%, funded by a $500k AFR loan from parents; S-corp from Year 2 (ownership thread's recommendation)", 2 + N + 2)
r += 1
d = {}
d["bal"] = srow("Loan balance, start of year", lambda c, y: "=LoanAmt" if y == 1 else f"=MAX(0,{L(3 + y - 1)}{r}-LoanForgive)", total=False)
d["int"] = srow("Interest paid to parents", lambda c, y: f"={c}{d['bal']}*LoanRate")
d["s"] = srow("S-corp year? (1 = yes)", lambda c, y: f"=IF({c}${SY}>=SCorpYear,1,0)", fmt=NUM, total=False)
d["sal"] = srow("Ryan's S-corp salary", lambda c, y: f"={c}{d['s']}*RyanSalary*(1+WageGrowth)^MAX(0,{c}${SY}-SCorpYear)")
d["inc"] = srow("Business income before owner pay (EBITDA + Ryan's hours - fed depreciation - interest)", lambda c, y: f"={c}{T['ebitda']}+{c}{T['ownlab']}-{c}{T['fdep']}-{c}{d['int']}")
d["emp"] = srow("Employer share of payroll tax on salary (deductible)", lambda c, y: f"={c}{d['sal']}*(SCorpPayroll-0.0765)")
d["se"] = srow("SE tax (LLC year) or payroll tax on salary (S-corp years)",
               lambda c, y: f"=IF({c}{d['s']}=1,{c}{d['sal']}*SCorpPayroll,IF({c}{d['inc']}>0,MIN({c}{d['inc']}*0.9235,SSBase*(1+WageGrowth)^({c}${SY}-1))*0.124+{c}{d['inc']}*0.9235*0.029,0))")
d["qbi"] = srow("QBI deduction (on profit after salary)", lambda c, y: f"=MAX(0,{c}{d['inc']}-IF({c}{d['s']}=1,{c}{d['sal']}+{c}{d['emp']},{c}{d['se']}/2))*QBIRate")
d["ti"] = srow("Federal taxable income (salary + K-1)", lambda c, y: f"={c}{d['inc']}-IF({c}{d['s']}=1,{c}{d['emp']},{c}{d['se']}/2)-{c}{d['qbi']}")
fd = carry_block("Fed", d["ti"], "MIN(RyanOtherInc,EBLSingle)", "RyanFed")
d["njti"] = srow("NJ taxable income", lambda c, y: f"={c}{T['ebitda']}+{c}{T['ownlab']}-{c}{T['ndep']}-{c}{d['int']}-{c}{d['emp']}")
nd = carry_block("NJ", d["njti"], "0", "RyanNJ")
d["tax"] = srow("Total tax (Ryan: federal + SE/payroll + NJ)", lambda c, y: f"={c}{fd['tax']}+{c}{d['se']}+{c}{nd['tax']}", bold=True)
d["pint"] = srow("Parents' after-tax interest received (stays in the family)", lambda c, y: f"={c}{d['int']}*(1-ParentFed-NIIT-ParentNJ)")
d["fam"] = srow("Family after-tax cash flow", lambda c, y: f"={c}{T['cf']}+{c}{T['ownlab']}-{c}{d['tax']}-{c}{d['int']}+{c}{d['pint']}-{c}{T['exit']}*ExitTax", bold=True, y0=f"=C{T['cf']}")
d["ryan"] = srow("Ryan's own after-tax cash flow (after interest to parents)", lambda c, y: f"={c}{T['cf']}+{c}{T['ownlab']}-{c}{d['tax']}-{c}{d['int']}-{c}{T['exit']}*ExitTax", bold=True, y0=f"=C{T['cf']}+LoanAmt")
blocks["D"] = d
r += 1

section(st, r, "Comparison", 6)
r += 1
for j, h in enumerate(["Structure", "NPV @ discount rate", "After-tax IRR", "Total tax yrs 1-17", "Year-1 tax (neg = saving)"]):
    c = st.cell(r, 2 + j, h)
    c.font = BOLD
    c.fill = SFILL
r += 1
COMP = {}
desc = {"A": "A  Ryan owns (gift), single-member LLC", "B": "B  Parents own, passive, no passive income", "C": "C  Parents own, passive, with passive income",
        "D": "D  Ryan owns, AFR loan from parents, S-corp yr 2+"}
for k in ("A", "B", "C", "D"):
    b = blocks[k]
    st.cell(r, 2, desc[k]).font = BLACK
    st.cell(r, 3, f"=C{b['fam']}+NPV(DiscRate,D{b['fam']}:{ycol(N)}{b['fam']})").number_format = USD
    st.cell(r, 4, f"=IRR(C{b['fam']}:{ycol(N)}{b['fam']})").number_format = PCT
    st.cell(r, 5, f"={L(3 + N + 1)}{b['tax']}").number_format = USD
    st.cell(r, 6, f"=D{b['tax']}").number_format = USD
    COMP[k] = r
    r += 1
st.cell(r + 1, 2, "Simplifications: marginal-rate method; NOLs used without the 80% limit; Ryan's QBI assumed below the threshold; SE tax ignores Ryan's other wages; "
        "exit taxed at one blended rate. Have a CPA run the final structure.").font = NOTE
st.freeze_panes = "C6"

# ---------------------------------------------------------------- Seller recast
rc = wb.create_sheet("Seller Recast", 1)
setw(rc, {"A": 2, "B": 44, "C": 14, "D": 14, "E": 14, "F": 80})
title(rc, "Broker P&L vs. underwritten Year 1 (annual)",
      "Broker column = proforma P&L dated 10-4-26 x 12 (photo IMG_1029). Underwritten = Projection Year 1.")
r = 5
for j, h in enumerate(["Line", "Broker P&L", "Underwritten Yr 1", "Difference", "Why"]):
    c = rc.cell(r, 2 + j, h)
    c.font = HDR
    c.fill = HFILL
r += 1
RC = [
    ("Self-serve", 216000, "ss", "Same; verify with CCI reports."),
    ("Wash-dry-fold", 120000, "wf", "Same in Year 1; price increase starts Year 2."),
    ("Ancillary (dry cleaning, ATM, vending, games)", 40200, "anc", "Haircut on unverifiable cash lines."),
    ("TOTAL REVENUE", 376200, "rev", "Listing says $372k; P&L sums to $376k."),
    ("Rent incl. CAM & sewer", 33096, "rent", ""),
    ("Payroll wages", 78084, "wages", "P&L: 91 hrs x $16.50. 2027 NJ minimum wage is $16.48, so $17.50 to keep staff."),
    ("Payroll taxes, workers' comp, sick leave", 0, "burden", "P&L shows N/A. Legally required."),
    ("Water", 32400, "water", "Plus 10% contingency until bills are seen."),
    ("PSE&G gas + electric", 32400, "gas", "Electric marked 'Inc'; confirm with bills."),
    ("Repairs & maintenance", 5400, "rep", "15-yr fleet; seller likely does own repairs unpaid."),
    ("Insurance", 2880, "ins", "$240/mo is too low for property + GL + equipment breakdown."),
    ("Card processing", 4800, "card", ""),
    ("POS system & cards", 6600, "pos", ""),
    ("W&F supplies", 3600, "wfs", ""),
    ("Vending COGS + dry cleaning cost", 10500, "cogs", ""),
    ("Phone, alarm, trash", 2220, "tel", "Underwritten adds the $375/yr fire alarm noted on the P&L."),
    ("Misc. supplies", 1800, "misc", ""),
    ("Accounting, tax & legal", 0, "pro", "Not in P&L."),
    ("Licenses & fees", 0, "lic", "Not in P&L."),
    ("Payroll service", 0, "psvc", "Not in P&L."),
    ("Marketing", 0, "mkt", "Not in P&L."),
    ("Contingency", 0, "cont", "1% of revenue."),
    ("TOTAL EXPENSES", None, "opex", "Broker's lines add to $17,815/mo; the P&L's TOTAL says $16,540."),
    ("NET OPERATING INCOME (absentee)", None, "ebitda", "Broker's lines, added correctly: $162,420, not $177,720."),
]
RCR = {}
for lbl, broker, key, why in RC:
    bold = lbl.isupper()
    rc.cell(r, 2, lbl).font = BOLD if bold else BLACK
    if broker is None:
        broker = f"=SUM(C{RCR['rev'] + 1}:C{r - 1})" if key == "opex" else f"=C{RCR['rev']}-C{RCR['opex']}"
    c = rc.cell(r, 3, broker)
    c.font = BLUE if not str(broker).startswith("=") else BLACK
    c.number_format = USD
    c = rc.cell(r, 4, f"=Projection!D{P[key]}")
    c.font = GREEN
    c.number_format = USD
    c = rc.cell(r, 5, f"=D{r}-C{r}")
    c.number_format = USD
    rc.cell(r, 6, why).font = NOTE
    if bold:
        for cc in range(2, 7):
            rc.cell(r, cc).border = TOP
    RCR[key] = r
    r += 1
r += 1
rc.cell(r, 2, "Broker's stated NOI (P&L TOTAL omits vending COGS, dry cleaning cost and card fees: $1,275/mo)").font = BLACK
c = rc.cell(r, 3, 177720)
c.font = BLUE
c.number_format = USD
RCR["stated"] = r
r += 1
rc.cell(r, 2, "Note: underwritten Year 1 uses the selected capex scenario; under scenarios 2/3 new machines lift revenue and cut utilities/repairs.").font = NOTE
r += 2
rc.cell(r, 2, "Implied price multiples").font = BOLD
r += 1
for lbl, f in (("Ask / broker's stated NOI", f"=Price/C{RCR['stated']}"),
               ("Ask / underwritten NOI (absentee)", f"=Price/D{RCR['ebitda']}"),
               ("Ask / underwritten SDE (Ryan works his hours)", "=Price/Projection!D%d" % P["sde"]),
               ("Ask / underwritten revenue", f"=Price/D{RCR['rev']}")):
    rc.cell(r, 2, lbl).font = BLACK
    c = rc.cell(r, 4, f)
    c.number_format = MULT
    RCR[lbl] = r
    r += 1
r += 1
rc.cell(r, 2, "Underwritten NOI after equipment reserve (avg. annual capex)").font = BOLD
c = rc.cell(r, 4, f"=D{RCR['ebitda']}-Projection!C{P['avgcapex']}")
c.number_format = USD
RCR["noi_res"] = r

# ---------------------------------------------------------------- Valuation
va = wb.create_sheet("Valuation", 2)
setw(va, {"A": 2, "B": 52, "C": 14, "D": 14, "E": 14, "F": 14, "G": 50})
title(va, "Valuation: what is this store worth?", "Three approaches. Market multiples come from the Comps sheet and research/comps.md.")
r = 5
section(va, r, "Earnings bases (Year 1, underwritten)", 6)
r += 1
V = {}
for k, lbl, f in (("noi", "NOI, absentee (all hours paid)", f"='Seller Recast'!D{RCR['ebitda']}"),
                  ("sde", "SDE, Ryan works his hours", f"=Projection!D{P['sde']}"),
                  ("res", "NOI after equipment reserve", f"='Seller Recast'!D{RCR['noi_res']}"),
                  ("broker", "Broker's stated NOI", f"='Seller Recast'!C{RCR['stated']}")):
    va.cell(r, 2, lbl).font = BLACK
    c = va.cell(r, 3, f)
    c.font = GREEN
    c.number_format = USD
    V[k] = r
    r += 1
r += 1
section(va, r, "1. Market approach: earnings x comparable multiples", 6)
r += 1
mults = [("Low (old fleet, NYC metro sold median less)", 2.5), ("Sold-listing median, NYC metro", 3.3), ("National closed-deal median (BizBuySell 2021-25)", 3.5), ("Rule of thumb, 15+ yr lease, low rent", 4.0)]
for j, h in enumerate(["Multiple", "x", "x NOI", "x SDE", "x NOI after reserve"]):
    c = va.cell(r, 2 + j, h)
    c.font = BOLD
    c.fill = SFILL
r += 1
mstart = r
for lbl, m in mults:
    va.cell(r, 2, lbl).font = BLACK
    c = va.cell(r, 3, m)
    c.font = BLUE
    c.number_format = MULT
    for j, k in enumerate(("noi", "sde", "res")):
        c = va.cell(r, 4 + j, f"=$C{r}*$C${V[k]}")
        c.number_format = USD
    r += 1
va.cell(mstart, 7, "Sources: research/comps.md (BizBuySell Insight 855 deals; 28 NYC-metro listings; Laundromat Resource rules of thumb)").font = NOTE
r += 1
section(va, r, "2. Income approach: price that still earns the target return", 6)
r += 1
va.cell(r, 2, "Max price at pre-tax target return (absentee cash flows, selected capex)").font = BLACK
pv = f"NPV(PreTaxTarget,Projection!D{P['cf']}:Projection!{ycol(N)}{P['cf']})"
c = va.cell(r, 3, f"=({pv}+Projection!C{P['capex']}-ClosingCosts-WorkingCap)")
c.number_format = USD
V["maxpre"] = r
va.cell(r, 7, "PV of Years 1-17 at the target, less Year-0 capex, closing costs and working capital.").font = NOTE
r += 1
for k in ("A", "B", "C", "D"):
    va.cell(r, 2, f"Max price at after-tax discount rate: structure {k}").font = BLACK
    c = va.cell(r, 3, f"=Price+Structures!C{COMP[k]}")
    c.number_format = USD
    V["max" + k] = r
    va.cell(r, 7, "Ask + NPV (approx.; taxes held at the current price's allocation).").font = NOTE
    r += 1
r += 1
section(va, r, "3. Asset-light floor", 6)
r += 1
va.cell(r, 2, "Used-equipment value + improvements (liquidation-ish floor)").font = BLACK
c = va.cell(r, 3, "=AllocEquip+AllocLHI")
c.number_format = USD
r += 2
section(va, r, "Conclusion", 6)
r += 1
va.cell(r, 2, "Fair value range: 3.0x (low) to 3.5x (high) underwritten absentee NOI").font = BOLD
c = va.cell(r, 3, f"=3.0*C{V['noi']}")
c.number_format = USD
c = va.cell(r, 4, f"=3.5*C{V['noi']}")
c.number_format = USD
V["range"] = r
r += 1
va.cell(r, 2, "Asking price").font = BLACK
va.cell(r, 3, "=Price").number_format = USD
r += 1
va.cell(r, 2, "Suggested opening offer (3.0x absentee NOI, rounded)").font = BOLD
c = va.cell(r, 3, f"=ROUND(3.0*C{V['noi']},-4)")
c.number_format = USD
V["offer"] = r
r += 1
va.cell(r, 2, "Walk-away price (3.5x absentee NOI, rounded)").font = BOLD
c = va.cell(r, 3, f"=ROUND(3.5*C{V['noi']},-4)")
c.number_format = USD
V["walk"] = r

# ---------------------------------------------------------------- Comps
cp = wb.create_sheet("Comps")
setw(cp, {"A": 2, "B": 34, "C": 12, "D": 12, "E": 12, "F": 9, "G": 9, "H": 30, "I": 10, "J": 60})
title(cp, "Comparable laundromat listings and sales", "Asking prices unless noted. 'Sold*' = broker marked sold; closing price not disclosed. Full table: /mnt/project-files/underwriting/research/comps.md")
r = 5
heads = ["Location", "Price", "Gross", "Cash flow", "P/CF", "P/Gross", "Notes", "Status", "Source"]
for j, h in enumerate(heads):
    c = cp.cell(r, 2 + j, h)
    c.font = HDR
    c.fill = HFILL
r += 1
comps = [
    ("Jersey City (SUBJECT, broker figures)", 450000, 372000, 177000, "~2,000 sf; 33W/36D ~15 yr; 17 yr lease, $2,758/mo", "Asking", "https://www.bizbuysell.com/business-opportunity/profitable-laundromat-run-absentee-high-revenue-jersey-city-nj-450k/2564518/"),
    ("Hudson County, NJ", 550000, 294000, 102000, "48 machines, energy-efficient", "Sold*", "https://www.hedgestone.com/business-opportunity/63086/profitable-absentee-run-laundromat-for-sale-in-new-jersey/"),
    ("Hudson County, NJ (downtown)", 50000, 104000, 43000, "15W/16D; ~5.5 yr lease left", "Sold*", "https://www.vestedbb.com/very-successful-laundromat-for-sale-in-hudson-county-new-jersey-listing-id-7091"),
    ("Bayonne, NJ", 179000, 175000, 16000, "Search-page figures only", "Asking", "https://www.bizbuysell.com/business-opportunity/laundromat-in-great-location-excellent-potential-bayonne-nj-179-000/2556056/"),
    ("Bergen County, NJ (plaza)", 950000, 498922, 173422, "4,000 sf; 50W/52D 3-7 yr; lease to 2041; rent 24%", "Asking", "https://www.bizbuysell.com/business-opportunity/4-000-sf-laundromat-for-sale-busy-shopping-plaza-with-parking/2521410/"),
    ("Union County, NJ", 650000, 416000, 150000, "2,800 sf; 41W/46D", "Sold*", "https://www.hedgestone.com/business-opportunity/62794/coin-operated-new-jersey-laundromat-for-sale/"),
    ("Bergen County, NJ", 110000, 153000, 58045, "", "Sold*", "https://www.vestedbb.com/great-laundromat-for-sale-in-bergen-county-new-jersey-listing-id-7167"),
    ("Ridgefield Park, NJ (24-hr)", 275000, None, 52000, "Unattended", "Asking", "https://www.bizbuysell.com/business-opportunity/stress-free-absentee-laundromat-open-24-hours/2560849/"),
    ("Paterson, NJ", 295000, 300000, None, "20-25 yr equipment; '2x net'", "Sold*", "https://www.laundromat-broker.com/listings.html"),
    ("South NJ (near Waterford)", 297500, 190000, 90000, "19 mo lease left", "Sold* 2025", "https://www.laundromat-broker.com/listings.html"),
    ("Bronx, NY (same broker template)", 450000, 400000, 177000, "32W/34D 11-12 yr; 3 yr lease", "Asking", "https://www.bizbuysell.com/business-opportunity/absentee-and-high-volume-laundromat-excellent-profit-bronx-ny-450k/2564395/"),
    ("Park Slope, Brooklyn", 799000, 467000, 156000, "Newer equipment, long lease", "Asking", "https://www.bizbuysell.com/new-york/new-york-city-metro-area/laundromats-and-coin-laundry-businesses-for-sale/"),
    ("Queens (prime)", 589999, 420000, 130000, "64 machines; new 10-yr lease", "Asking", "https://www.bizbuysell.com/new-york/new-york-city-metro-area/laundromats-and-coin-laundry-businesses-for-sale/"),
    ("Prospect Park 11226", 449300, None, 160840, "75 machines; 12 yr lease", "Asking", "https://www.bizbuysell.com/new-york/new-york-city-metro-area/laundromats-and-coin-laundry-businesses-for-sale/"),
    ("Brooklyn 11219", 349500, None, 112950, "41 machines; 15-yr lease option", "Asking", "https://www.bizbuysell.com/new-york/new-york-city-metro-area/laundromats-and-coin-laundry-businesses-for-sale/"),
]
cstart = r
for loc, pr, gr, cf, note, status, src in comps:
    cp.cell(r, 2, loc).font = BOLD if "SUBJECT" in loc else BLACK
    for j, v in enumerate((pr, gr, cf)):
        c = cp.cell(r, 3 + j, v if v is not None else "n/a")
        c.font = BLUE
        c.number_format = USD
    cp.cell(r, 6, f'=IF(ISNUMBER(E{r}),C{r}/E{r},"n/a")').number_format = MULT
    cp.cell(r, 7, f'=IF(ISNUMBER(D{r}),C{r}/D{r},"n/a")').number_format = MULT
    cp.cell(r, 8, note).font = NOTE
    cp.cell(r, 9, status).font = BLACK
    cp.cell(r, 10, src).font = NOTE
    r += 1
cend = r - 1
cp.cell(r, 2, "Median of comps (excl. subject)").font = BOLD
cp.cell(r, 6, f"=MEDIAN(F{cstart+1}:F{cend})").number_format = MULT
cp.cell(r, 7, f"=MEDIAN(G{cstart+1}:G{cend})").number_format = MULT
r += 2
bench = [
    ("BizBuySell national closed deals 2021-25 (n=855): median price/SDE", "3.50x (IQR 2.72-4.50x); 2025 avg 4.12x"),
    ("BizBuySell national: median price/revenue", "1.21x; sale price ~92% of ask"),
    ("NYC-metro current asking (n=28): median price/cash flow", "3.5x (range 1.45-7.0x)"),
    ("Laundromat Resource rule of thumb: 13-16 yr equipment", "3.5-4.0x NOI"),
    ("Laundromat Resource rule of thumb: 15+ yr lease; rent <=25% of gross", "4.5-5.0x NOI"),
    ("Utilities, % of gross (2023 operator survey)", "21.6% avg; ~15% with new machines"),
    ("BizBuySell EDGE, NJ laundromat listings: revenue (25th / median / 75th pct)", "$172K / $312K / $495K (listing PDF, 10/14/2026)"),
    ("BizBuySell EDGE, NJ laundromat listings: SDE (25th / median / 75th pct)", "$62.5K / $107K / $246K"),
    ("BizBuySell EDGE, operating NJ laundromats 2024 averages", "Revenue $570K; SDE $127K; EBITDA $81K (14% of revenue)"),
]
cp.cell(r, 2, "Benchmarks").font = BOLD
r += 1
for a_, b_ in bench:
    cp.cell(r, 2, a_).font = BLACK
    cp.cell(r, 8, b_).font = BLUE
    r += 1

# ---------------------------------------------------------------- Summary (first sheet)
sm = wb.create_sheet("Summary", 0)
setw(sm, {"A": 2, "B": 58, "C": 16, "D": 16, "E": 16, "F": 16})
title(sm, "Tumble Wash Laundromat: underwriting summary",
      "1147 Summit Ave, Jersey City NJ. BizBuySell #2564518. Model as of 2026-10-16. Change assumptions on the Inputs sheet.")
r = 5
section(sm, r, "Deal", 5)
r += 1
for lbl, f, fmt in (("Asking price", "=Price", USD), ("Total cash at close (price + closing + working capital)", "=TotalInvest", USD),
                    ("Capex scenario selected (1 keep/repair, 2 washers now, 3 full retool)", "=CapexScenario", NUM),
                    ("Equipment capex over the lease (inflated)", f"=-SUM(Projection!C{P['capex']}:Projection!{ycol(N)}{P['capex']})", USD)):
    sm.cell(r, 2, lbl).font = BLACK
    c = sm.cell(r, 3, f)
    c.font = GREEN
    c.number_format = fmt
    r += 1
r += 1
section(sm, r, "Earnings: broker vs. underwritten (Year 1)", 5)
r += 1
for j, h in enumerate(["", "Broker", "Underwritten", "Difference"]):
    sm.cell(r, 2 + j, h).font = BOLD
r += 1
for lbl, key in (("Revenue", "rev"), ("Operating expenses", "opex"), ("NOI, absentee", "ebitda")):
    sm.cell(r, 2, lbl).font = BLACK
    for j, col in enumerate("CDE"):
        c = sm.cell(r, 3 + j, f"='Seller Recast'!{col}{RCR[key]}")
        c.font = GREEN
        c.number_format = USD
    r += 1
sm.cell(r, 2, "Broker's stated NOI (its P&L total leaves out $15,300/yr of its own expense lines)").font = BLACK
c = sm.cell(r, 3, f"='Seller Recast'!C{RCR['stated']}")
c.number_format = USD
c.font = GREEN
r += 1
sm.cell(r, 2, "SDE if Ryan works his hours").font = BLACK
c = sm.cell(r, 4, f"=Projection!D{P['sde']}")
c.number_format = USD
c.font = GREEN
r += 1
sm.cell(r, 2, "NOI after average equipment reserve").font = BLACK
c = sm.cell(r, 4, f"='Seller Recast'!D{RCR['noi_res']}")
c.number_format = USD
c.font = GREEN
r += 1
sm.cell(r, 2, "Ask as a multiple of underwritten absentee NOI").font = BLACK
c = sm.cell(r, 4, f"=Price/'Seller Recast'!D{RCR['ebitda']}")
c.number_format = MULT
r += 2
section(sm, r, "Returns", 5)
r += 1
sm.cell(r, 2, "Pre-tax unlevered IRR, 17 years (absentee)").font = BLACK
c = sm.cell(r, 3, f"=Projection!C{P['irr']}")
c.number_format = PCT
c.font = GREEN
r += 1
sm.cell(r, 2, "Year-1 cash-on-cash, pre-tax").font = BLACK
c = sm.cell(r, 3, f"=Projection!C{P['coc']}")
c.number_format = PCT
c.font = GREEN
r += 1
for j, h in enumerate(["After tax, family level", "NPV", "IRR", "Total tax 1-17", "Year-1 tax"]):
    sm.cell(r, 2 + j, h).font = BOLD
r += 1
for k in ("A", "B", "C", "D"):
    sm.cell(r, 2, desc[k]).font = BLACK
    for j, col in enumerate("CDEF"):
        c = sm.cell(r, 3 + j, f"=Structures!{col}{COMP[k]}")
        c.font = GREEN
        c.number_format = PCT if col == "D" else USD
    r += 1
r += 1
section(sm, r, "Value", 5)
r += 1
for lbl, f in (("Fair value range, low", f"=Valuation!C{V['range']}"), ("Fair value range, high", f"=Valuation!D{V['range']}"),
               ("Max price at pre-tax target return", f"=Valuation!C{V['maxpre']}"), ("Suggested opening offer", f"=Valuation!C{V['offer']}"), ("Walk-away price", f"=Valuation!C{V['walk']}")):
    sm.cell(r, 2, lbl).font = BLACK
    c = sm.cell(r, 3, f)
    c.font = GREEN
    c.number_format = USD
    r += 1
r += 1
SNAP_ROW = r
section(sm, r, "Scenario snapshot (values pasted by build script; rerun build_model.py after changing inputs)", 5)

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
    for row_ in ws.iter_rows():
        for c in row_:
            if c.font is None or c.font.name != FONT:
                f = c.font
                c.font = Font(name=FONT, size=f.size or 10, bold=f.bold, italic=f.italic, color=f.color)

wb.save(OUT)
print(OUT, "snap_row", SNAP_ROW)
import json
json.dump({"snap_row": SNAP_ROW, "irr_row": P["irr"], "comp": COMP, "noi_row": RCR["ebitda"], "capex_row": P["capex"],
           "range": V["range"], "maxpre": V["maxpre"], "offer": V["offer"], "walk": V["walk"], "coc": P["coc"], "sde": P["sde"], "ebitda": P["ebitda"], "ownhrs": P["ownhrs"]}, open(OUT + ".meta.json", "w"))
