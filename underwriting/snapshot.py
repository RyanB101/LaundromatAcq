"""Runs scenario/sensitivity cases through LibreOffice and pastes a static table onto the Summary sheet.
Run after build_model.py:  python3 snapshot.py
"""
import json, os, shutil, subprocess, sys, tempfile
from openpyxl import load_workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
XL = os.path.join(HERE, "Tumble_Wash_Underwriting_Model.xlsx")
META = json.load(open(XL + ".meta.json"))
RECALC = sys.argv[1] if len(sys.argv) > 1 else "recalc.py"

CASES = [
    ("Base: ask $450k, keep & repair machines", {}),
    ("Offer $400k", {"Price": 400000}),
    ("Offer $350k", {"Price": 350000}),
    ("Replace washers at close (scenario 2)", {"CapexScenario": 2}),
    ("Full retool at close (scenario 3)", {"CapexScenario": 3}),
    ("Electric billed separately (+$12k/yr)", {"ElecExtra": 12000}),
    ("Revenue 10% below the P&L", {"SelfServe": 194400, "WFLbsMonth": 7200}),
    ("Ryan works 0 hrs (fully absentee)", {"OwnerHours": 0}),
    ("No lease renewal (exit value $0)", {"ExitMult": 0}),
    ("Your plan: you 22 hrs/wk Yr 1, ~10 after; day manager from Yr 2", {"OwnerHours": 22, "OwnerHoursLater": 10, "MgrHours": 40}),
    ("Your plan + staffed delivery/commercial (5,000 lb/mo)", {"OwnerHours": 22, "OwnerHoursLater": 10, "MgrHours": 40, "DelivLbsRamp": 2000, "DelivLbsFull": 5000}),
    ("BEAR case (cases.py)", {"OwnerHours": 22, "OwnerHoursLater": 10, "MgrHours": 40, "SelfServe": 194400, "SSGrowth": 0.02, "WFLbsMonth": 7200, "WFPriceNew": 1.40, "AncRealize": 0.7, "ElecExtra": 12000, "RepairsYr1": 15000, "CapexScenario": 4, "RetoolLift": 0.0, "ExitMult": 1.5}),
    ("BASE case (cases.py)", {"OwnerHours": 22, "OwnerHoursLater": 10, "MgrHours": 40}),
    ("BULL case (cases.py)", {"OwnerHours": 22, "OwnerHoursLater": 10, "MgrHours": 40, "UtilCont": 0, "AncRealize": 1.0, "WFPriceNew": 1.65, "SelfServe": 226800, "SSGrowth": 0.035, "RepairsYr1": 8000, "DelivLbsRamp": 1500, "DelivLbsFull": 4000, "RetoolLift": 0.10, "ExitMult": 2.5}),
    ("Broker's expense assumptions", {"Wage": 16.5, "Burden": 0, "Insurance": 2880, "RepairsYr1": 5400, "UtilCont": 0,
                                      "AncRealize": 1, "ProFees": 0, "Licenses": 0, "PayrollSvc": 0, "Marketing": 0, "ContPct": 0}),
]


def refs(wb):
    out = {}
    for n, dn in wb.defined_names.items():
        sheet, cell = dn.attr_text.split("!")
        out[n] = (sheet, cell.replace("$", ""))
    return out


def run_case(over):
    wb = load_workbook(XL)
    rf = refs(wb)
    for k, v in over.items():
        sh, cell = rf[k]
        wb[sh][cell] = v
    tmp = tempfile.mkdtemp()
    p = os.path.join(tmp, "case.xlsx")
    wb.save(p)
    subprocess.run([sys.executable, RECALC, p, "90"], check=True, capture_output=True)
    v = load_workbook(p, data_only=True)
    sr, pj, stc = v["Seller Recast"], v["Projection"], v["Structures"]
    res = {
        "noi": sr[f"D{META['noi_row']}"].value,
        "irr": pj[f"C{META['irr_row']}"].value,
        "npvA": stc[f"C{META['comp']['A']}"].value,
        "irrA": stc[f"D{META['comp']['A']}"].value,
        "npvB": stc[f"C{META['comp']['B']}"].value,
        "npvC": stc[f"C{META['comp']['C']}"].value,
        "npvD": stc[f"C{META['comp']['D']}"].value,
    }
    shutil.rmtree(tmp)
    return res


rows = [(lbl, run_case(o)) for lbl, o in CASES]
wb = load_workbook(XL)
sm = wb["Summary"]
r = META["snap_row"] + 1
heads = ["Case", "Yr-1 NOI", "Pre-tax IRR", "A: NPV", "A: IRR", "B: NPV", "C: NPV", "D: NPV"]
for j, h in enumerate(heads):
    sm.cell(r, 2 + j, h).font = Font(name="Arial", bold=True, size=10)
r += 1
fmts = ['$#,##0;($#,##0)', '0.0%', '$#,##0;($#,##0)', '0.0%', '$#,##0;($#,##0)', '$#,##0;($#,##0)', '$#,##0;($#,##0)']
for lbl, res in rows:
    sm.cell(r, 2, lbl).font = Font(name="Arial", size=10)
    for j, k in enumerate(("noi", "irr", "npvA", "irrA", "npvB", "npvC", "npvD")):
        c = sm.cell(r, 3 + j, res[k])
        c.number_format = fmts[j]
        c.font = Font(name="Arial", size=10)
    r += 1
sm.cell(r, 2, "Static values from snapshot.py; each case changes only the listed inputs from the base. NPV at the Inputs discount rate.").font = Font(name="Arial", italic=True, size=9, color="595959")
sm.column_dimensions["G"].width = 14
sm.column_dimensions["H"].width = 14
sm.column_dimensions["I"].width = 14
wb.save(XL)
subprocess.run([sys.executable, RECALC, XL, "90"], check=True)
for lbl, res in rows:
    print(lbl, {k: (round(v, 3) if isinstance(v, float) else v) for k, v in res.items()})
